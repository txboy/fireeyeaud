"""把用户模板语法预处理成 docxtpl（Jinja2）标记并渲染。

用户语法（P2 已固化）：

- 顶层字段 ``{{项目名称}}``：Jinja2 支持中文标识符，原样保留。
- 循环字段 ``{{疑点.标题}}`` → ``{{ finding["标题"] }}``；``{{取证单.标题}}`` → ``{{ form["标题"] }}``。
- 段落循环 ``{{#疑点}}`` / ``{{/疑点}}`` → ``{%p for finding in findings %}`` / ``{%p endfor %}``。
- 表格行循环 → ``{%tr for ... %}`` / ``{%tr endfor %}``（标记所在段落属于表格单元格时）。

渲染本身只读模板、只写目标 docx，不访问数据库。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph

from ..errors import ApiError

_LOOP_START = re.compile(r"^\s*\{\{#\s*([^{}]+?)\s*\}\}\s*$")
_LOOP_END = re.compile(r"^\s*\{\{/\s*([^{}]+?)\s*\}\}\s*$")

# 循环名 → (Jinja 变量名, 上下文集合名)
LOOP_VARS: dict[str, tuple[str, str]] = {
    "疑点": ("finding", "findings"),
    "取证单": ("form", "forms"),
}

FIELD_IN_LOOP: dict[str, re.Pattern[str]] = {
    "疑点": re.compile(r"\{\{\s*疑点\.([^{}]+?)\s*\}\}"),
    "取证单": re.compile(r"\{\{\s*取证单\.([^{}]+?)\s*\}\}"),
}


def _replace_in_paragraph(paragraph: Paragraph, text: str) -> None:
    """把整段文本写入首个 run（其余清空），避免占位符被拆到多个 run 里。"""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _in_table_cell(paragraph: Paragraph) -> bool:
    """段落是否位于表格单元格内（决定用 tr 还是 p 循环标记）。"""
    parent = paragraph._p.getparent()
    return bool(parent is not None and parent.tag.endswith("tc"))


def _process_paragraph(paragraph: Paragraph, kind: str) -> None:
    text = "".join(run.text for run in paragraph.runs)
    if not text:
        return

    start = _LOOP_START.match(text)
    end = _LOOP_END.match(text)
    if start or end:
        match = start or end
        assert match is not None
        name = match.group(1).strip()
        if name not in LOOP_VARS:
            raise ApiError(
                "E_TEMPLATE_LOOP", 400, "error.templateLoopUnsupported", {"name": name}
            )
        var, coll = LOOP_VARS[name]
        tag = "tr" if _in_table_cell(paragraph) else "p"
        if start:
            marker = "{%" + tag + f" for {var} in {coll}" + " %}"
        else:
            marker = "{%" + tag + " endfor %}"
        _replace_in_paragraph(paragraph, marker)
        return

    new_text = text
    for name, (var, _coll) in LOOP_VARS.items():
        pattern = FIELD_IN_LOOP[name]
        new_text = pattern.sub(lambda m, v=var: '{{ ' + v + '["' + m.group(1) + '"] }}', new_text)
    if new_text != text:
        _replace_in_paragraph(paragraph, new_text)


def _process_paragraphs(paragraphs: list[Paragraph], kind: str) -> None:
    for paragraph in paragraphs:
        _process_paragraph(paragraph, kind)


def _process_tables(document: Any, kind: str) -> None:
    for table in document.tables:
        _process_table(table, kind)


def _find_loop_start(row: Any) -> tuple[Paragraph, str] | None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            text = "".join(run.text for run in paragraph.runs)
            match = _LOOP_START.match(text)
            if match:
                return paragraph, match.group(1).strip()
    return None


def _clear_marker_paragraphs(row: Any) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            text = "".join(run.text for run in paragraph.runs)
            if _LOOP_START.match(text) or _LOOP_END.match(text):
                _replace_in_paragraph(paragraph, "")


def _process_table(table: Any, kind: str) -> None:
    """处理表格。

    docxtpl 的 ``{%tr for %}`` 会**丢弃标记所在行**，循环体是标记行与 endfor 行之间的行。
    P2 的写法把开始标记放在数据行的第一个单元格里，因此这里克隆该行：
    原行改成标记行，克隆行作为重复的循环体。
    """
    import copy

    from docx.table import _Row

    for row in list(table.rows):
        marker = _find_loop_start(row)
        if marker is None:
            for cell in row.cells:
                _process_paragraphs(cell.paragraphs, kind)
                for nested in cell.tables:
                    _process_table(nested, kind)
            continue

        marker_paragraph, name = marker
        if name not in LOOP_VARS:
            raise ApiError(
                "E_TEMPLATE_LOOP", 400, "error.templateLoopUnsupported", {"name": name}
            )
        var, coll = LOOP_VARS[name]

        clone_tr = copy.deepcopy(row._tr)
        row._tr.addnext(clone_tr)
        clone_row = _Row(clone_tr, table)

        # 原行：整行会被 docxtpl 丢弃，只保留循环标记
        _replace_in_paragraph(marker_paragraph, "{%tr " + f"for {var} in {coll}" + " %}")
        # 克隆行：作为循环体，去掉标记、转换字段
        _clear_marker_paragraphs(clone_row)
        for cell in clone_row.cells:
            _process_paragraphs(cell.paragraphs, kind)


def prepare_template(src: Path, dst: Path, kind: str) -> None:
    """把用户模板转成 docxtpl 可渲染的中间文件。"""
    doc = Document(str(src))
    _process_paragraphs(doc.paragraphs, kind)
    _process_tables(doc, kind)
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            _process_paragraphs(part.paragraphs, kind)
            # 页眉/页脚里的表格同样要走行循环克隆，否则数据行会被 docxtpl 丢掉。
            _process_tables(part, kind)
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))


def _jinja_env() -> Any:
    """沙箱 + 自动转义。

    - 沙箱：模板来自用户上传的 docx，不能让 `{{ ''.__class__.__mro__ }}` 之类执行任意代码。
    - 自动转义：渲染结果直接写进 docx 的 XML，`&` / `<` 不转义会破坏文档结构
      （实测 `A&B` 会截断后续内容，`<...>` 甚至能把表格整段吃掉）。
    """
    from jinja2.sandbox import SandboxedEnvironment

    return SandboxedEnvironment(autoescape=True)


def render_docx(template: Path, out: Path, context: dict[str, Any]) -> None:
    """用 docxtpl 渲染到 ``out``；失败统一转 ApiError。"""
    from docxtpl import DocxTemplate

    try:
        tpl = DocxTemplate(str(template))
        tpl.render(context, jinja_env=_jinja_env())
        out.parent.mkdir(parents=True, exist_ok=True)
        tpl.save(str(out))
    except ApiError:
        raise
    except Exception as exc:  # 模板语法/渲染错误统一报模板错误
        raise ApiError(
            "E_TEMPLATE_RENDER", 400, "error.templateRender", {"reason": str(exc)}
        ) from exc