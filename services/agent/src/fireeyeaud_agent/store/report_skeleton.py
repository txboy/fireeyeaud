"""报告模板骨架解析（方案 A）。

从 docx 模板里读出：页面设置、标题样式、各级标题样式、正文样式、表格样式、
固定段落与章节骨架。导出时用这些样式把结构化正文写回模板，模型不接触格式。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from ..errors import ApiError

_HEADING_RE = re.compile(r"(?:标题|Heading|head)\s*([1-9])", re.IGNORECASE)
_ALIGN = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def _walk(style: Any, getter: Any) -> Any:
    """沿样式继承链取值（样式自身 → base_style）。"""
    current = style
    while current is not None:
        value = getter(current)
        if value is not None:
            return value
        current = getattr(current, "base_style", None)
    return None


def _font_attr(style: Any, name: str) -> Any:
    return _walk(style, lambda item: getattr(item.font, name, None))


def _pf_attr(style: Any, name: str) -> Any:
    return _walk(style, lambda item: getattr(item.paragraph_format, name, None))


def _east_asia(element: Any) -> str | None:
    rpr = getattr(element, "rPr", None)
    if rpr is None:
        return None
    fonts = rpr.rFonts
    if fonts is None:
        return None
    return fonts.get(qn("w:eastAsia"))


def _outline_level(paragraph: Any) -> int | None:
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.outlineLvl is not None:
        try:
            return int(ppr.outlineLvl.val) + 1
        except (TypeError, ValueError):
            return None
    style_ppr = getattr(paragraph.style.element, "pPr", None)
    if style_ppr is not None and style_ppr.outlineLvl is not None:
        try:
            return int(style_ppr.outlineLvl.val) + 1
        except (TypeError, ValueError):
            return None
    return None


def _style_of(paragraph: Any) -> dict[str, Any]:
    style = paragraph.style
    run = paragraph.runs[0] if paragraph.runs else None
    fmt = paragraph.paragraph_format

    size = (run.font.size if run is not None else None) or _font_attr(style, "size")
    font = (run.font.name if run is not None else None) or _font_attr(style, "name")
    east = (
        (_east_asia(run._element) if run is not None else None)
        or _walk(style, lambda item: _east_asia(item.element))
    )
    bold = (run.font.bold if run is not None else None)
    if bold is None:
        bold = _font_attr(style, "bold")

    line = fmt.line_spacing
    if line is None:
        line = _pf_attr(style, "line_spacing")
    rule = fmt.line_spacing_rule
    if rule is None:
        rule = _pf_attr(style, "line_spacing_rule")
    indent = fmt.first_line_indent
    if indent is None:
        indent = _pf_attr(style, "first_line_indent")
    before = fmt.space_before
    if before is None:
        before = _pf_attr(style, "space_before")
    after = fmt.space_after
    if after is None:
        after = _pf_attr(style, "space_after")
    align = paragraph.alignment
    if align is None:
        align = _pf_attr(style, "alignment")

    return {
        "styleName": style.name,
        "font": font,
        "eastAsia": east,
        "size": round(size.pt, 2) if size is not None else None,
        "bold": bool(bold) if bold is not None else None,
        "align": _ALIGN.get(align, "left") if align is not None else None,
        "lineSpacing": round(line.pt, 2) if hasattr(line, "pt") else line,
        "lineRule": str(rule).split(" ")[0] if rule is not None else None,
        "firstLineIndent": round(indent.pt, 2) if indent is not None else None,
        "spaceBefore": round(before.pt, 2) if before is not None else None,
        "spaceAfter": round(after.pt, 2) if after is not None else None,
    }


def _level_of(paragraph: Any, text: str, body_size: float | None) -> int | None:
    match = _HEADING_RE.search(paragraph.style.name or "")
    if match:
        return int(match.group(1))
    outline = _outline_level(paragraph)
    if outline:
        return outline
    # 回退：加粗且明显大于正文、且不是长段落
    if not paragraph.runs or len(text) > 40:
        return None
    run = paragraph.runs[0]
    size = run.font.size.pt if run.font.size is not None else None
    if size is None:
        style_size = _font_attr(paragraph.style, "size")
        size = style_size.pt if style_size is not None else None
    bold = run.font.bold
    if bold is None:
        bold = _font_attr(paragraph.style, "bold")
    if bold and size and body_size and size >= body_size + 2:
        return 1
    return None


def _mode_size(paragraphs: list[Any]) -> float | None:
    counts: dict[float, int] = {}
    for paragraph in paragraphs:
        text = (paragraph.text or "").strip()
        if len(text) < 8:
            continue
        run = paragraph.runs[0] if paragraph.runs else None
        size = run.font.size if run is not None else None
        if size is None:
            size = _font_attr(paragraph.style, "size")
        if size is None:
            continue
        key = round(size.pt, 1)
        counts[key] = counts.get(key, 0) + len(text)
    if not counts:
        return None
    return max(counts.items(), key=lambda item: item[1])[0]


def template_text(path: Path) -> dict[str, Any]:
    """整份模板原文（段落 + 表格，按文档顺序），作为给模型看的参考范本。"""
    from docx.table import Table

    document = Document(str(Path(path)))
    blocks: list[dict[str, Any]] = []
    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            text = "".join(node.text or "" for node in element.iter(qn("w:t")))
            if text.strip():
                blocks.append({"kind": "paragraph", "text": text})
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            rows = [
                [
                    "".join(node.text or "" for node in cell._tc.iter(qn("w:t")))
                    for cell in row.cells
                ]
                for row in table.rows
            ]
            blocks.append({"kind": "table", "rows": rows})
    return {"blocks": blocks}


def analyze_template(path: Path) -> dict[str, Any]:
    path = Path(path)
    if not path.is_file():
        raise ApiError(
            "E_FILE_NOT_FOUND", 404, "error.fileNotFound", {"name": path.name}
        )
    if path.suffix.lower() != ".docx":
        raise ApiError("E_FILE_TYPE", 400, "error.fileType", {"name": path.name})
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ApiError(
            "E_FILE_TYPE", 400, "error.fileType", {"name": path.name}
        ) from exc

    warnings: list[str] = []
    paragraphs = list(document.paragraphs)
    filled = [p for p in paragraphs if (p.text or "").strip()]

    body_size = _mode_size(filled)

    title_paragraph = None
    for paragraph in filled:
        style = _style_of(paragraph)
        if (
            style["align"] == "center"
            and style["size"]
            and style["size"] >= 20
            and len(paragraph.text.strip()) <= 60
        ):
            title_paragraph = paragraph
            break
    if title_paragraph is None and filled:
        first = filled[0]
        size = _style_of(first)["size"] or 0
        if len(first.text.strip()) <= 40 and size and body_size and size > body_size:
            title_paragraph = first

    title = (
        {"text": title_paragraph.text.strip(), "style": _style_of(title_paragraph)}
        if title_paragraph is not None
        else None
    )

    headings: dict[int, dict[str, Any]] = {}
    sections: list[dict[str, Any]] = []
    preamble: list[str] = []
    seen_heading = False
    current: dict[str, Any] | None = None

    for paragraph in filled:
        if paragraph is title_paragraph:
            continue
        text = paragraph.text.strip()
        level = _level_of(paragraph, text, body_size)
        if level:
            seen_heading = True
            headings.setdefault(level, {"level": level, "style": _style_of(paragraph)})
            current = {
                "title": text,
                "level": level,
                "sampleText": [],
                "style": _style_of(paragraph),
            }
            sections.append(current)
            continue
        if current is not None:
            if len(current["sampleText"]) < 3:
                current["sampleText"].append(text)
        elif not seen_heading:
            preamble.append(text)

    if not sections:
        warnings.append("未识别到章节标题：将按通用报告结构生成，导出时使用正文样式。")
    if title_paragraph is None:
        warnings.append("未识别到独立标题段：导出的报告标题将使用正文样式加粗居中。")
    # 占位符不再需要填值，模板只作参考范本；这里不再给出警告，避免误导。

    body_source = None
    for paragraph in filled:
        if paragraph is title_paragraph:
            continue
        if _level_of(paragraph, paragraph.text.strip(), body_size):
            continue
        if len(paragraph.text.strip()) >= 8:
            body_source = paragraph
            break
    if body_source is None:
        body_source = filled[-1] if filled else None
    body_style = _style_of(body_source) if body_source is not None else {}

    tables: list[dict[str, Any]] = []
    for table in document.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        first_cell = table.rows[0].cells[0]
        sample = first_cell.paragraphs[0] if first_cell.paragraphs else None
        tables.append(
            {
                "columns": len(table.columns),
                "rows": len(table.rows),
                "headers": headers,
                "styleName": table.style.name if table.style is not None else None,
                "cellStyle": _style_of(sample) if sample is not None else {},
                "headerBold": True,
            }
        )

    section = document.sections[0] if document.sections else None
    page = {
        "widthCm": round(section.page_width.cm, 2) if section and section.page_width else None,
        "heightCm": round(section.page_height.cm, 2) if section and section.page_height else None,
        "marginLeftCm": round(section.left_margin.cm, 2) if section and section.left_margin else None,
        "marginRightCm": round(section.right_margin.cm, 2) if section and section.right_margin else None,
        "marginTopCm": round(section.top_margin.cm, 2) if section and section.top_margin else None,
        "marginBottomCm": round(section.bottom_margin.cm, 2) if section and section.bottom_margin else None,
    }

    confidence = "high" if sections and title is not None else ("medium" if sections else "low")
    return {
        "schemaVersion": 1,
        "source": path.name,
        "page": page,
        "title": title,
        "body": body_style,
        "headings": [headings[key] for key in sorted(headings)],
        "sections": sections,
        "tables": tables,
        "preamble": preamble,
        "templateText": template_text(path),
        "warnings": warnings,
        "confidence": confidence,
    }
