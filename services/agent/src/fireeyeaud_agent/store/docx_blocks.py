"""docx 块级渲染原语：正文块归一化、摘要规范化与按骨架样式写入。

报告（``report_docs``）与取证单（``form_docs``）共用；本模块不依赖任何 store 模块。
"""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

_MAX_BLOCKS = 60
_MAX_TEXT = 20000
_MAX_TABLE_ROWS = 200

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


# ---- 摘要 / 大纲归一化 ----


def normalize_outline(
    sections: list[Any], *, existing: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    by_id = {str(item.get("id")): item for item in existing}
    used: set[str] = set()
    out: list[dict[str, Any]] = []
    counter = 1
    for raw in sections or []:
        if not isinstance(raw, dict):
            continue
        title = str(raw.get("title") or "").strip()[:200]
        if not title:
            continue
        try:
            level = int(raw.get("level") or 1)
        except (TypeError, ValueError):
            level = 1
        level = min(max(level, 1), 3)
        section_id = str(raw.get("id") or "")
        if not section_id or section_id in used:
            while f"S{counter}" in used:
                counter += 1
            section_id = f"S{counter}"
        used.add(section_id)
        suffix = section_id[1:]
        if section_id.startswith("S") and suffix.isdigit():
            counter = max(counter, int(suffix) + 1)
        previous = by_id.get(section_id) or {}
        out.append(
            {
                "id": section_id,
                "title": title,
                "level": level,
                "goal": str(raw.get("goal") or previous.get("goal") or "").strip()[:500],
                "blocks": list(previous.get("blocks") or []),
            }
        )
        if len(out) >= 30:
            break
    return out


# ---- 正文块归一化 ----


def _clean_refs(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    refs = [str(item).strip() for item in value if str(item).strip()]
    return list(dict.fromkeys(refs))[:20]


def _normalize_runs(value: Any) -> list[dict[str, Any]] | None:
    """规范化行内格式；返回 None 表示没有可用的富文本信息。"""
    if not isinstance(value, list):
        return None
    runs: list[dict[str, Any]] = []
    for raw in value:
        if not isinstance(raw, dict):
            continue
        text = str(raw.get("text") or "")[:_MAX_TEXT]
        if not text:
            continue
        run: dict[str, Any] = {"text": text}
        for key in ("bold", "italic", "underline"):
            if raw.get(key):
                run[key] = True
        runs.append(run)
    return runs or None


def _runs_text(runs: list[dict[str, Any]]) -> str:
    return "".join(str(run.get("text") or "") for run in runs).strip()[:_MAX_TEXT]


def _normalize_align(value: Any) -> str | None:
    align = str(value or "")
    return align if align in {"left", "center", "right", "justify"} else None


def normalize_blocks(blocks: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for raw in (blocks or [])[:_MAX_BLOCKS]:
        if not isinstance(raw, dict):
            continue
        kind = str(raw.get("type") or "paragraph")
        refs = _clean_refs(raw.get("sourceRefs"))
        origin = str(raw.get("origin") or "form")
        if origin not in {"form", "user", "model"}:
            origin = "form"
        if kind == "paragraph":
            runs = _normalize_runs(raw.get("runs"))
            text = (
                _runs_text(runs)
                if runs
                else str(raw.get("text") or "").strip()[:_MAX_TEXT]
            )
            if text:
                block: dict[str, Any] = {
                    "type": "paragraph",
                    "text": text,
                    "sourceRefs": refs,
                    "origin": origin,
                }
                if runs:
                    block["runs"] = runs
                align = _normalize_align(raw.get("align"))
                if align:
                    block["align"] = align
                out.append(block)
        elif kind == "bullets":
            raw_items = raw.get("items") if isinstance(raw.get("items"), list) else []
            raw_item_runs = (
                raw.get("itemRuns") if isinstance(raw.get("itemRuns"), list) else []
            )
            raw_item_align = (
                raw.get("itemAlign") if isinstance(raw.get("itemAlign"), list) else []
            )
            count = max(len(raw_items), len(raw_item_runs))
            items: list[str] = []
            item_runs: list[list[dict[str, Any]]] = []
            item_align: list[str | None] = []
            for index in range(count):
                runs = (
                    _normalize_runs(raw_item_runs[index])
                    if index < len(raw_item_runs)
                    else None
                )
                if runs:
                    text = _runs_text(runs)
                else:
                    text = str(
                        raw_items[index] if index < len(raw_items) else ""
                    ).strip()[:_MAX_TEXT]
                if not text:
                    continue
                if len(items) >= 50:
                    break
                items.append(text)
                item_runs.append(runs or [])
                item_align.append(
                    _normalize_align(raw_item_align[index])
                    if index < len(raw_item_align)
                    else None
                )
            if items:
                block = {
                    "type": "bullets",
                    "items": items,
                    "sourceRefs": refs,
                    "origin": origin,
                }
                if any(item_runs):
                    block["itemRuns"] = item_runs
                if any(item_align):
                    block["itemAlign"] = item_align
                out.append(block)
        elif kind == "table":
            raw_columns = (
                raw.get("columns") if isinstance(raw.get("columns"), list) else []
            )
            raw_column_runs = (
                raw.get("columnRuns")
                if isinstance(raw.get("columnRuns"), list)
                else []
            )
            column_count = max(len(raw_columns), len(raw_column_runs))
            raw_column_align = (
                raw.get("columnAlign")
                if isinstance(raw.get("columnAlign"), list)
                else []
            )
            columns: list[str] = []
            column_runs: list[list[dict[str, Any]]] = []
            column_align: list[str | None] = []
            for index in range(min(column_count, 12)):
                runs = (
                    _normalize_runs(raw_column_runs[index])
                    if index < len(raw_column_runs)
                    else None
                )
                text = (
                    _runs_text(runs)
                    if runs
                    else str(
                        raw_columns[index] if index < len(raw_columns) else ""
                    ).strip()
                )
                columns.append(text)
                column_runs.append(runs or [])
                column_align.append(
                    _normalize_align(raw_column_align[index])
                    if index < len(raw_column_align)
                    else None
                )

            raw_rows = raw.get("rows") if isinstance(raw.get("rows"), list) else []
            raw_row_runs = (
                raw.get("rowRuns") if isinstance(raw.get("rowRuns"), list) else []
            )
            raw_row_align = (
                raw.get("rowAlign") if isinstance(raw.get("rowAlign"), list) else []
            )
            rows: list[list[str]] = []
            row_runs: list[list[list[dict[str, Any]]]] = []
            row_align: list[list[str | None]] = []
            for row_index, raw_row in enumerate(raw_rows[:_MAX_TABLE_ROWS]):
                cells = raw_row if isinstance(raw_row, list) else []
                raw_cell_runs = (
                    raw_row_runs[row_index]
                    if row_index < len(raw_row_runs)
                    and isinstance(raw_row_runs[row_index], list)
                    else []
                )
                raw_cell_align = (
                    raw_row_align[row_index]
                    if row_index < len(raw_row_align)
                    and isinstance(raw_row_align[row_index], list)
                    else []
                )
                row: list[str] = []
                cell_runs: list[list[dict[str, Any]]] = []
                cell_align: list[str | None] = []
                for column_index in range(len(columns)):
                    runs = (
                        _normalize_runs(raw_cell_runs[column_index])
                        if column_index < len(raw_cell_runs)
                        else None
                    )
                    text = (
                        _runs_text(runs)
                        if runs
                        else str(
                            cells[column_index] if column_index < len(cells) else ""
                        ).strip()[:2000]
                    )
                    row.append(text)
                    cell_runs.append(runs or [])
                    cell_align.append(
                        _normalize_align(raw_cell_align[column_index])
                        if column_index < len(raw_cell_align)
                        else None
                    )
                rows.append(row)
                row_runs.append(cell_runs)
                row_align.append(cell_align)
            if columns:
                block = {
                    "type": "table",
                    "columns": columns,
                    "rows": rows,
                    "sourceRefs": refs,
                    "origin": origin,
                }
                if any(column_runs):
                    block["columnRuns"] = column_runs
                if any(any(cell for cell in row) for row in row_runs):
                    block["rowRuns"] = row_runs
                if any(column_align):
                    block["columnAlign"] = column_align
                if any(any(cell for cell in row) for row in row_align):
                    block["rowAlign"] = row_align
                out.append(block)
    return out


# ---- 按骨架样式写入 ----


def apply_style(paragraph: Any, style: dict[str, Any], document: Any) -> None:
    name = str(style.get("styleName") or "")
    if name:
        try:
            paragraph.style = document.styles[name]
        except KeyError:
            pass

    pf = paragraph.paragraph_format
    align = _ALIGN.get(str(style.get("align") or ""))
    if align is not None:
        paragraph.alignment = align

    spacing = style.get("lineSpacing")
    if isinstance(spacing, (int, float)):
        rule = str(style.get("lineRule") or "").upper()
        if rule == "EXACTLY":
            pf.line_spacing = Pt(float(spacing))
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        elif rule == "AT_LEAST":
            pf.line_spacing = Pt(float(spacing))
            pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        else:
            # MULTIPLE 存的是倍数；没有规则时按固定磅值处理更安全
            pf.line_spacing = (
                float(spacing) if rule == "MULTIPLE" else Pt(float(spacing))
            )

    for key, setter in (
        ("firstLineIndent", "first_line_indent"),
        ("spaceBefore", "space_before"),
        ("spaceAfter", "space_after"),
    ):
        value = style.get(key)
        if isinstance(value, (int, float)):
            setattr(pf, setter, Pt(float(value)))


def _apply_run(run: Any, style: dict[str, Any]) -> None:
    if style.get("font"):
        run.font.name = str(style["font"])
    east = style.get("eastAsia")
    if east:
        rpr = run._element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), str(east))
    size = style.get("size")
    if isinstance(size, (int, float)):
        run.font.size = Pt(float(size))
    if style.get("bold") is not None:
        run.font.bold = bool(style["bold"])


def _add_runs(
    paragraph: Any,
    text: str,
    style: dict[str, Any],
    runs: list[dict[str, Any]] | None = None,
) -> None:
    if runs:
        for item in runs:
            value = str(item.get("text") or "")
            if not value:
                continue
            run = paragraph.add_run(value)
            _apply_run(run, style)
            if item.get("bold"):
                run.font.bold = True
            if item.get("italic"):
                run.font.italic = True
            if item.get("underline"):
                run.font.underline = True
        return
    run = paragraph.add_run(text)
    _apply_run(run, style)


def _apply_align(paragraph: Any, align: Any) -> None:
    value = _ALIGN.get(str(align or ""))
    if value is not None:
        paragraph.alignment = value


def add_text(
    document: Any,
    text: str,
    style: dict[str, Any],
    *,
    bullet: bool = False,
    runs: list[dict[str, Any]] | None = None,
    align: str | None = None,
) -> Any:
    paragraph = document.add_paragraph()
    apply_style(paragraph, style, document)
    _apply_align(paragraph, align)
    if bullet:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        prefix = paragraph.add_run("· ")
        _apply_run(prefix, style)
    _add_runs(paragraph, text, style, runs)
    return paragraph


def write_cell(
    cell: Any,
    text: str,
    style: dict[str, Any],
    document: Any,
    runs: list[dict[str, Any]] | None = None,
    align: str | None = None,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    apply_style(paragraph, style, document)
    _apply_align(paragraph, align)
    _add_runs(paragraph, text, style, runs)


def _element_text(element: Any) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def clear_body(document: Any, keep_until: set[str]) -> list[str]:
    """保留标题之前的固定内容（红头 / 封面表格），清空标题之后。

    返回被清掉的非空固定文字，供导出时提示用户哪些内容不会保留。
    """
    body = document.element.body
    children = [
        child for child in list(body.iterchildren()) if not child.tag.endswith("}sectPr")
    ]
    title_index = None
    for index, child in enumerate(children):
        if child.tag.endswith("}p") and _element_text(child) in keep_until:
            title_index = index
            break

    if title_index is None:
        lost = [text for text in (_element_text(c) for c in children) if text]
        for child in children:
            body.remove(child)
        return lost

    lost: list[str] = []
    for child in children[title_index:]:
        text = _element_text(child)
        if text:
            lost.append(text)
        body.remove(child)
    return lost


def heading_style(skeleton: dict[str, Any], level: int) -> dict[str, Any]:
    headings = skeleton.get("headings") or []
    for item in headings:
        if int(item.get("level") or 0) == level:
            return dict(item.get("style") or {})
    if headings:
        base = dict(headings[0].get("style") or {})
        body = dict(skeleton.get("body") or {})
        base_size = float(base.get("size") or body.get("size") or 16)
        base["size"] = base_size
        base["bold"] = True
        return base
    body = dict(skeleton.get("body") or {})
    return {
        **body,
        "bold": True,
        "size": float(body.get("size") or 16) + 2,
        "firstLineIndent": 0,
    }