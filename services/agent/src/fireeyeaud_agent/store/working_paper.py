"""审计工作底稿生成与产物清单（设计 §14.15）。

结构：封面 / 审计结论 / 疑点汇总 / 证据索引 / 附录。
产物：`.fireeye/artifacts/working-paper/<id>/{底稿.*, manifest.json}`。
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from ..errors import ApiError
from ..util import now_iso
from .evidence import load_evidence
from .findings import load_findings
from .project import load_project, next_counter

_TITLE = {"zh-CN": "审计工作底稿", "en-US": "Audit Working Paper"}
_HEADINGS = {
    "zh-CN": {
        "conclusion": "审计结论",
        "findings": "疑点汇总",
        "evidence": "证据索引",
        "appendix": "附录",
        "workspace": "工作区",
        "generatedAt": "生成时间",
        "counts": "疑点统计",
        "id": "编号",
        "severity": "风险",
        "status": "状态",
        "title": "标题",
        "amount": "金额",
        "subjects": "科目",
        "tool": "工具",
        "step": "步骤",
        "ts": "时间",
    },
    "en-US": {
        "conclusion": "Conclusion",
        "findings": "Findings summary",
        "evidence": "Evidence index",
        "appendix": "Appendix",
        "workspace": "Workspace",
        "generatedAt": "Generated at",
        "counts": "Findings",
        "id": "ID",
        "severity": "Severity",
        "status": "Status",
        "title": "Title",
        "amount": "Amount",
        "subjects": "Subjects",
        "tool": "Tool",
        "step": "Step",
        "ts": "Time",
    },
}
_SEVERITY = {
    "zh-CN": {"high": "高", "medium": "中", "low": "低"},
    "en-US": {"high": "High", "medium": "Medium", "low": "Low"},
}
_STATUS = {
    "zh-CN": {
        "pending": "待复核",
        "in_review": "复核中",
        "confirmed": "已确认",
        "rejected": "已驳回",
        "false_positive": "误报",
    },
    "en-US": {
        "pending": "Pending",
        "in_review": "In review",
        "confirmed": "Confirmed",
        "rejected": "Rejected",
        "false_positive": "False positive",
    },
}
_FILE = {"word": "底稿.docx", "xlsx": "底稿.xlsx", "pdf": "底稿.pdf"}


def working_paper_dir(root: Path) -> Path:
    return Path(root) / ".fireeye" / "artifacts" / "working-paper"


def collect_material(root: Path, finding_ids: list[str] | None) -> dict[str, Any]:
    root = Path(root)
    fe = root / ".fireeye"
    findings = load_findings(root)
    if finding_ids:
        wanted = set(finding_ids)
        findings = [f for f in findings if f.get("id") in wanted]
    evidence_by_id = {e["id"]: e for e in load_evidence(root, limit=100000)}
    project = load_project(fe) or {}
    return {
        "workspaceName": project.get("name") or root.name,
        "generatedAt": now_iso(),
        "findings": findings,
        "evidence": evidence_by_id,
    }


def _labels(locale: str) -> dict[str, str]:
    return _HEADINGS.get(locale, _HEADINGS["zh-CN"])


def _title(locale: str) -> str:
    return _TITLE.get(locale, _TITLE["zh-CN"])


def _finding_title(finding: dict[str, Any]) -> str:
    title = finding.get("title")
    return (
        (title or {}).get("text") or finding.get("ruleKey") or finding.get("id") or ""
    )


def _counts(findings: list[dict[str, Any]]) -> dict[str, int]:
    result = {"high": 0, "medium": 0, "low": 0}
    for f in findings:
        key = f.get("severity")
        if key in result:
            result[key] += 1
    return result


def _write_manifest(base: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    base.mkdir(parents=True, exist_ok=True)
    (base / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), "utf-8"
    )
    return manifest


def _render_docx(base: Path, material: dict[str, Any], locale: str) -> None:
    from docx import Document

    L = _labels(locale)
    findings = material["findings"]
    doc = Document()
    doc.add_heading(_title(locale), level=0)
    doc.add_paragraph(f"{L['workspace']}: {material['workspaceName']}")
    doc.add_paragraph(f"{L['generatedAt']}: {material['generatedAt']}")

    doc.add_heading(L["conclusion"], level=1)
    counts = _counts(findings)
    doc.add_paragraph(
        f"{L['counts']}: " + ", ".join(f"{k}={v}" for k, v in counts.items())
    )

    doc.add_heading(L["findings"], level=1)
    table = doc.add_table(rows=1, cols=6)
    for i, key in enumerate(
        ["id", "severity", "status", "title", "amount", "subjects"]
    ):
        table.rows[0].cells[i].text = L[key]
    sev = _SEVERITY.get(locale, _SEVERITY["zh-CN"])
    st = _STATUS.get(locale, _STATUS["zh-CN"])
    for f in findings:
        cells = table.add_row().cells
        cells[0].text = f.get("id", "")
        cells[1].text = sev.get(f.get("severity"), f.get("severity", ""))
        cells[2].text = st.get(f.get("status"), f.get("status", ""))
        cells[3].text = _finding_title(f)
        cells[4].text = str(f.get("amount") or "")
        cells[5].text = "、".join(f.get("subjects") or [])

    doc.add_heading(L["evidence"], level=1)
    for f in findings:
        for ev_id in f.get("evidenceRefs") or []:
            ev = material["evidence"].get(ev_id)
            if not ev:
                continue
            doc.add_paragraph(
                f"{f.get('id')} · {ev.get('tool') or ev.get('step')} · {ev.get('ts')}"
            )

    doc.add_heading(L["appendix"], level=1)
    doc.add_paragraph("templateId: default")
    doc.save(str(base / _FILE["word"]))


def _render_xlsx(base: Path, material: dict[str, Any], locale: str) -> None:
    from openpyxl import Workbook

    L = _labels(locale)
    sev = _SEVERITY.get(locale, _SEVERITY["zh-CN"])
    st = _STATUS.get(locale, _STATUS["zh-CN"])
    findings = material["findings"]
    wb = Workbook()

    summary = wb.active
    summary.title = L["findings"]
    summary.append(
        [L["id"], L["severity"], L["status"], L["title"], L["amount"], L["subjects"]]
    )
    for f in findings:
        summary.append(
            [
                f.get("id", ""),
                sev.get(f.get("severity"), f.get("severity", "")),
                st.get(f.get("status"), f.get("status", "")),
                _finding_title(f),
                str(f.get("amount") or ""),
                "、".join(f.get("subjects") or []),
            ]
        )

    evidence = wb.create_sheet(L["evidence"])
    evidence.append([L["id"], L["tool"], L["step"], L["ts"]])
    for f in findings:
        for ev_id in f.get("evidenceRefs") or []:
            ev = material["evidence"].get(ev_id)
            if not ev:
                continue
            evidence.append(
                [
                    f.get("id", ""),
                    ev.get("tool") or "",
                    ev.get("step") or "",
                    ev.get("ts") or "",
                ]
            )
    wb.save(str(base / _FILE["xlsx"]))


def _render_pdf(base: Path, material: dict[str, Any], locale: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    L = _labels(locale)
    findings = material["findings"]
    sev = _SEVERITY.get(locale, _SEVERITY["zh-CN"])
    st = _STATUS.get(locale, _STATUS["zh-CN"])
    title_style = ParagraphStyle("title", fontName=font, fontSize=18, leading=24)
    head_style = ParagraphStyle("head", fontName=font, fontSize=13, leading=18)
    body_style = ParagraphStyle("body", fontName=font, fontSize=10, leading=14)

    story: list[Any] = [
        Paragraph(_title(locale), title_style),
        Spacer(1, 6 * mm),
        Paragraph(f"{L['workspace']}: {material['workspaceName']}", body_style),
        Paragraph(f"{L['generatedAt']}: {material['generatedAt']}", body_style),
        Spacer(1, 4 * mm),
        Paragraph(L["conclusion"], head_style),
    ]
    counts = _counts(findings)
    story.append(
        Paragraph(
            f"{L['counts']}: " + ", ".join(f"{k}={v}" for k, v in counts.items()),
            body_style,
        )
    )
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(L["findings"], head_style))

    data = [[L["id"], L["severity"], L["status"], L["title"], L["amount"]]]
    for f in findings:
        data.append(
            [
                f.get("id", ""),
                sev.get(f.get("severity"), f.get("severity", "")),
                st.get(f.get("status"), f.get("status", "")),
                _finding_title(f),
                str(f.get("amount") or ""),
            ]
        )
    table = Table(data, colWidths=[22 * mm, 18 * mm, 24 * mm, 70 * mm, 30 * mm])
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
            ]
        )
    )
    story.append(table)

    doc = SimpleDocTemplate(str(base / _FILE["pdf"]), pagesize=A4)
    doc.build(story)


def create_working_paper(
    root: Path,
    *,
    format: str,
    locale: str = "zh-CN",
    finding_ids: list[str] | None = None,
    template_id: str = "default",
) -> dict[str, Any]:
    root = Path(root)
    if format not in _FILE:
        raise ApiError("E_EXPORT_FORMAT", 400, "error.exportFormat", {"format": format})
    fe = root / ".fireeye"
    material = collect_material(root, finding_ids)
    seq = next_counter(fe, "workingPaper", root)
    wp_id = f"wp_{seq:04d}"
    base = working_paper_dir(root) / wp_id
    base.mkdir(parents=True, exist_ok=True)

    if format == "word":
        _render_docx(base, material, locale)
    elif format == "xlsx":
        _render_xlsx(base, material, locale)
    elif format == "pdf":
        _render_pdf(base, material, locale)

    manifest = {
        "id": wp_id,
        "createdAt": material["generatedAt"],
        "format": format,
        "locale": locale,
        "templateId": template_id,
        "findingIds": [f["id"] for f in material["findings"]],
        "artifactRef": f"working-paper/{wp_id}",
        "files": [_FILE[format]],
    }
    return _write_manifest(base, manifest)


def list_working_papers(root: Path) -> list[dict[str, Any]]:
    base = working_paper_dir(Path(root))
    if not base.exists():
        return []
    items: list[dict[str, Any]] = []
    for child in base.iterdir():
        manifest = child / "manifest.json"
        if not manifest.exists():
            continue
        try:
            data = json.loads(manifest.read_text("utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        items.append(data)
    items.sort(key=lambda d: (d.get("createdAt", ""), d.get("id", "")), reverse=True)
    return items
