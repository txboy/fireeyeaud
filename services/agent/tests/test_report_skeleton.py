"""模板骨架解析：页面、标题/正文字体字号、章节、表格。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from fireeyeaud_agent.store.report_skeleton import analyze_template, template_text


def _set_east_asia(style_or_run, name: str) -> None:
    element = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rpr = element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def _build_template(path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)

    normal = doc.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(16)
    _set_east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = Pt(30)
    normal.paragraph_format.first_line_indent = Pt(32)

    heading = doc.styles.add_style("标题 1", 1)  # WD_STYLE_TYPE.PARAGRAPH
    heading.base_style = normal
    heading.font.name = "黑体"
    heading.font.size = Pt(16)
    heading.font.bold = True
    _set_east_asia(heading, "黑体")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("某某项目审计报告")
    run.font.name = "小标宋"
    run.font.size = Pt(22)
    _set_east_asia(run, "小标宋")

    doc.add_paragraph("根据年度审计计划，现开展专项审计。")

    doc.add_paragraph("一、审计概况", style=heading)
    doc.add_paragraph("本次审计范围包括财务收支与合同管理。")
    doc.add_paragraph("审计期间为2025年度。")

    doc.add_paragraph("二、审计发现", style=heading)
    doc.add_paragraph("发现的主要问题如下。")

    table = doc.add_table(rows=2, cols=3)
    table.style = doc.styles["Table Grid"]
    headers = ["序号", "疑点", "金额"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True
    for index, text in enumerate(["1", "示例疑点", "100元"]):
        table.rows[1].cells[index].text = text

    doc.add_paragraph("妥否，请批示。")
    doc.save(str(path))
    return path


def test_analyze_template(tmp_path):
    path = _build_template(tmp_path / "template.docx")
    skeleton = analyze_template(path)

    assert skeleton["page"]["widthCm"] == 21
    assert skeleton["page"]["marginLeftCm"] == 2.8

    assert skeleton["title"]["text"] == "某某项目审计报告"
    assert skeleton["title"]["style"]["size"] == 22
    assert skeleton["title"]["style"]["eastAsia"] == "小标宋"
    assert skeleton["title"]["style"]["align"] == "center"

    assert skeleton["body"]["size"] == 16
    assert skeleton["body"]["eastAsia"] == "仿宋_GB2312"
    assert round(skeleton["body"]["lineSpacing"]) == 30
    assert skeleton["body"]["firstLineIndent"] == 32

    assert [s["title"] for s in skeleton["sections"]] == ["一、审计概况", "二、审计发现"]
    assert skeleton["sections"][0]["level"] == 1
    assert any("财务收支" in text for text in skeleton["sections"][0]["sampleText"])

    assert skeleton["headings"][0]["level"] == 1
    assert skeleton["headings"][0]["style"]["eastAsia"] == "黑体"

    assert skeleton["tables"][0]["columns"] == 3
    assert skeleton["tables"][0]["headers"] == ["序号", "疑点", "金额"]

    assert skeleton["preamble"]
    assert skeleton["warnings"] == []


def test_template_without_headings(tmp_path):
    doc = Document()
    doc.add_paragraph("只有一段普通正文，没有任何标题样式。")
    path = tmp_path / "plain.docx"
    doc.save(str(path))

    skeleton = analyze_template(path)
    assert skeleton["sections"] == []
    assert any("未识别到章节标题" in w for w in skeleton["warnings"])


def test_template_text_keeps_paragraphs_tables_and_placeholders(tmp_path):
    doc = Document()
    doc.add_paragraph("项目：{{项目名称}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.标题}}"
    doc.add_paragraph("说明：{{#疑点}}…{{/疑点}}")
    doc.save(str(tmp_path / "t.docx"))

    text = template_text(tmp_path / "t.docx")
    assert text["blocks"][0] == {"kind": "paragraph", "text": "项目：{{项目名称}}"}
    assert text["blocks"][1]["kind"] == "table"
    assert text["blocks"][1]["rows"] == [["编号", "标题"], ["{{#疑点}}", "{{疑点.标题}}"]]
    assert text["blocks"][2]["text"] == "说明：{{#疑点}}…{{/疑点}}"


def test_template_text_skips_empty_paragraphs(tmp_path):
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    doc.add_paragraph("唯一一段")
    doc.save(str(tmp_path / "empty.docx"))
    assert template_text(tmp_path / "empty.docx")["blocks"] == [
        {"kind": "paragraph", "text": "唯一一段"}
    ]


def test_analyze_template_includes_template_text(tmp_path):
    path = _build_template(tmp_path / "template.docx")
    skeleton = analyze_template(path)
    assert skeleton["templateText"]["blocks"]
    assert any(
        block["kind"] == "table" and block["rows"][0] == ["序号", "疑点", "金额"]
        for block in skeleton["templateText"]["blocks"]
    )


def test_heading_falls_back_to_bold_larger_text(tmp_path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    doc.add_paragraph("普通正文内容，长度足够用于判断正文字号的大小。")

    heading = doc.add_paragraph()
    run = heading.add_run("第一部分 基本情况")
    run.bold = True
    run.font.size = Pt(15)
    doc.add_paragraph("这里是第一部分的内容。")

    path = tmp_path / "fallback.docx"
    doc.save(str(path))

    skeleton = analyze_template(path)
    assert [s["title"] for s in skeleton["sections"]] == ["第一部分 基本情况"]
    assert skeleton["sections"][0]["level"] == 1
