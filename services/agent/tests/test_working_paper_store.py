import json
from pathlib import Path

from fireeyeaud_agent.store.evidence import append_evidence
from fireeyeaud_agent.store.findings import create_finding
from fireeyeaud_agent.store.working_paper import (
    create_working_paper,
    list_working_papers,
)


def _seed(root: Path) -> dict:
    (root / ".fireeye").mkdir()
    ev = append_evidence(root, session_id="s_1", step="tool_call", tool="sql_query")
    return create_finding(
        root,
        title="大额往来",
        description="测试疑点",
        severity="high",
        dataRefs=[{"sourceId": "ds_1", "table": "src_ds1_序时账", "rowRef": "1182"}],
        sessionId="s_1",
        ruleKey="tie_check",
        evidenceRefs=[ev["id"]],
    )


def test_create_working_paper_docx(tmp_path):
    finding = _seed(tmp_path)
    manifest = create_working_paper(
        tmp_path, format="word", locale="zh-CN", finding_ids=None
    )
    assert manifest["id"].startswith("wp_")
    assert manifest["format"] == "word"
    assert manifest["findingIds"] == [finding["id"]]
    assert manifest["files"] == ["底稿.docx"]

    base = tmp_path / ".fireeye" / "artifacts" / "working-paper" / manifest["id"]
    assert (base / "manifest.json").exists()
    docx = base / "底稿.docx"
    assert docx.exists() and docx.stat().st_size > 0

    on_disk = json.loads((base / "manifest.json").read_text("utf-8"))
    assert on_disk["id"] == manifest["id"]


def test_create_working_paper_filter_by_finding(tmp_path):
    a = _seed(tmp_path)
    create_finding(
        tmp_path,
        title="b",
        description="d",
        severity="low",
        dataRefs=[],
        sessionId="s_1",
    )
    manifest = create_working_paper(
        tmp_path, format="word", locale="zh-CN", finding_ids=[a["id"]]
    )
    assert manifest["findingIds"] == [a["id"]]


def test_list_working_papers_newest_first(tmp_path):
    _seed(tmp_path)
    first = create_working_paper(
        tmp_path, format="word", locale="zh-CN", finding_ids=None
    )
    second = create_working_paper(
        tmp_path, format="word", locale="zh-CN", finding_ids=None
    )
    items = list_working_papers(tmp_path)
    assert [m["id"] for m in items] == [second["id"], first["id"]]


def test_list_working_papers_empty(tmp_path):
    (tmp_path / ".fireeye").mkdir()
    assert list_working_papers(tmp_path) == []


def test_create_working_paper_xlsx(tmp_path):
    finding = _seed(tmp_path)
    manifest = create_working_paper(
        tmp_path, format="xlsx", locale="zh-CN", finding_ids=None
    )
    assert manifest["files"] == ["底稿.xlsx"]
    base = tmp_path / ".fireeye" / "artifacts" / "working-paper" / manifest["id"]
    book = base / "底稿.xlsx"
    assert book.exists() and book.stat().st_size > 0
    from openpyxl import load_workbook

    wb = load_workbook(book)
    assert "疑点汇总" in wb.sheetnames
    ws = wb["疑点汇总"]
    assert ws.cell(row=1, column=1).value == "编号"
    assert ws.cell(row=2, column=1).value == finding["id"]


def test_create_working_paper_pdf(tmp_path):
    _seed(tmp_path)
    manifest = create_working_paper(
        tmp_path, format="pdf", locale="zh-CN", finding_ids=None
    )
    assert manifest["files"] == ["底稿.pdf"]
    base = tmp_path / ".fireeye" / "artifacts" / "working-paper" / manifest["id"]
    pdf = base / "底稿.pdf"
    assert pdf.exists() and pdf.stat().st_size > 0
    assert pdf.read_bytes()[:5] == b"%PDF-"


def test_create_working_paper_bad_format(tmp_path):
    (tmp_path / ".fireeye").mkdir()
    import pytest

    from fireeyeaud_agent.errors import ApiError

    with pytest.raises(ApiError):
        create_working_paper(tmp_path, format="doc", locale="zh-CN", finding_ids=None)
