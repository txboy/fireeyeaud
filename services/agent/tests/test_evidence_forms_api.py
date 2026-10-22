"""取证单 / 报告路由：预览 → 保存 → 列表 → 回执 → 删除。"""

from __future__ import annotations

from pathlib import Path

import httpx
from docx import Document

from fireeyeaud_agent.main import app
from fireeyeaud_agent.store.findings import create_finding
from fireeyeaud_agent.store.templates import upload_template


def _template(path: Path, kind: str) -> Path:
    doc = Document()
    if kind == "evidence-form":
        doc.add_paragraph("项目：{{项目名称}}")
        doc.add_paragraph("编号：{{取证单编号}}")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "编号"
        table.cell(0, 1).text = "标题"
        table.cell(1, 0).text = "{{#疑点}}"
        table.cell(1, 1).text = "{{疑点.标题}}"
        table.cell(2, 0).text = "{{/疑点}}"
    else:
        doc.add_paragraph("报告：{{报告编号}}")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "编号"
        table.cell(0, 1).text = "标题"
        table.cell(1, 0).text = "{{#取证单}}"
        table.cell(1, 1).text = "{{取证单.编号}} {{取证单.标题}}"
        table.cell(2, 0).text = "{{/取证单}}"
    doc.save(path)
    return path


async def _open(client, root: Path) -> None:
    await client.post("/api/v1/workspaces/open", json={"path": str(root)})


async def test_forms_flow(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    root = tmp_path / "ws"
    (root / ".fireeye").mkdir(parents=True)
    fid = create_finding(
        root,
        title="关联方交易",
        description="未披露",
        severity="high",
        dataRefs=[],
        sessionId="",
        origin="manual",
    )["id"]
    form_tpl = upload_template(
        root,
        _template(tmp_path / "f.docx", "evidence-form"),
        kind="evidence-form",
        scope="project",
    )
    report_tpl = upload_template(
        root,
        _template(tmp_path / "r.docx", "report"),
        kind="report",
        scope="project",
    )

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        await _open(client, root)

        preview = await client.post(
            "/api/v1/evidence-forms/preview",
            json={
                "templateId": form_tpl["id"],
                "templateScope": "project",
                "findingIds": [fid],
                "header": {"projectName": "年报审计"},
            },
        )
        assert preview.status_code == 200
        pv = preview.json()
        assert pv["relPath"].endswith("document.docx")

        committed = await client.post(
            "/api/v1/evidence-forms",
            json={"token": pv["token"], "title": "关联方取证单"},
        )
        assert committed.status_code == 201
        form = committed.json()
        assert form["id"].startswith("QZ-")
        assert form["templateScope"] == "project"

        listed = (await client.get("/api/v1/evidence-forms")).json()
        assert [i["id"] for i in listed["items"]] == [form["id"]]
        assert (await client.get(f"/api/v1/evidence-forms/{form['id']}")).status_code == 200

        receipt = tmp_path / "客户签收.pdf"
        receipt.write_bytes(b"%PDF")
        attached = await client.post(
            f"/api/v1/evidence-forms/{form['id']}/receipt",
            json={"path": str(receipt)},
        )
        assert attached.status_code == 200
        assert attached.json()["receipt"]["name"] == "客户签收.pdf"

        cleared = await client.delete(f"/api/v1/evidence-forms/{form['id']}/receipt")
        assert cleared.status_code == 200
        assert cleared.json()["receipt"] is None

        # 报告：用刚保存的取证单
        rpv = await client.post(
            "/api/v1/reports/preview",
            json={
                "templateId": report_tpl["id"],
                "templateScope": "project",
                "formIds": [form["id"]],
                "header": {"projectName": "年报审计"},
            },
        )
        assert rpv.status_code == 200
        report = (
            await client.post(
                "/api/v1/reports",
                json={"token": rpv.json()["token"], "title": "项目报告"},
            )
        ).json()
        assert report["id"].startswith("RPT-")
        assert report["formIds"] == [form["id"]]
        assert [i["id"] for i in (await client.get("/api/v1/reports")).json()["items"]] == [
            report["id"]
        ]

        assert (await client.delete(f"/api/v1/reports/{report['id']}")).status_code == 200
        assert (await client.delete(f"/api/v1/evidence-forms/{form['id']}")).status_code == 200
        assert (await client.get("/api/v1/evidence-forms")).json()["items"] == []
        await client.post("/api/v1/workspaces/close")


async def test_artifact_routes_require_workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        await client.post("/api/v1/workspaces/close")
        assert (await client.get("/api/v1/evidence-forms")).status_code == 409
        assert (await client.get("/api/v1/reports")).status_code == 409


async def test_artifact_404s(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    root = tmp_path / "ws"
    (root / ".fireeye").mkdir(parents=True)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        await _open(client, root)
        assert (await client.get("/api/v1/evidence-forms/QZ-9999")).status_code == 404
        assert (await client.delete("/api/v1/evidence-forms/QZ-9999")).status_code == 404
        assert (await client.get("/api/v1/reports/RPT-9999")).status_code == 404
        assert (
            await client.post("/api/v1/evidence-forms", json={"token": "nope"})
        ).status_code == 404
        await client.post("/api/v1/workspaces/close")