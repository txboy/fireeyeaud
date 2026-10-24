"""报告工作台 API：草稿、大纲、章节、校验、导出、对话记录。"""

from __future__ import annotations

import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from fireeyeaud_agent.config import store as config_store
from fireeyeaud_agent.main import app
from fireeyeaud_agent.store import evidence_forms
from fireeyeaud_agent.store import templates as templates_store


@pytest.fixture(autouse=True)
def _global_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))


def _east_asia(style_or_run, name: str) -> None:
    element = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _workspace(tmp_path: Path) -> tuple[Path, str]:
    root = tmp_path / "ws"
    (root / ".fireeye").mkdir(parents=True)
    (root / ".fireeye" / "project.json").write_text(
        json.dumps({"name": "示例项目", "counters": {}}, ensure_ascii=False), "utf-8"
    )

    doc = Document()
    doc.sections[0].page_width = Cm(21)
    doc.sections[0].page_height = Cm(29.7)
    normal = doc.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(16)
    _east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = Pt(30)
    normal.paragraph_format.first_line_indent = Pt(32)
    heading = doc.styles.add_style("标题 1", 1)
    heading.base_style = normal
    heading.font.name = "黑体"
    heading.font.size = Pt(16)
    heading.font.bold = True
    _east_asia(heading, "黑体")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("审计报告")
    run.font.size = Pt(22)
    run.font.name = "小标宋"
    _east_asia(run, "小标宋")
    doc.add_paragraph("一、基本情况", style=heading)
    doc.add_paragraph("占位内容。")

    item = config_store.create_item(
        "templates",
        {
            "scope": "project",
            "kind": "report",
            "name": "报告模板",
            "fileName": "report.docx",
        },
        root,
    )
    target = templates_store.templates_root("project", root) / "report" / "report.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))

    # 新格式取证单（写作台产物）：files 含 body.json，迁移不会清理它。
    form_dir = evidence_forms.forms_dir(root) / "QZ-0001"
    form_dir.mkdir(parents=True, exist_ok=True)
    evidence_forms.save_manifest(
        root,
        {
            "id": "QZ-0001",
            "kind": "evidence-form",
            "title": "采购取证单",
            "createdAt": "2026-10-22T10:00:00+08:00",
            "header": {"preparedBy": "张审计"},
            "receipt": None,
            "findingIds": ["F-0001"],
            "files": ["manifest.json", "body.json", "skeleton.json", "template.docx"],
            "findingsSnapshot": [
                {
                    "id": "F-0001",
                    "title": "服务器费用",
                    "description": "采购阿里云服务器，费用99元。",
                    "severity": "low",
                    "status": "confirmed",
                }
            ],
        },
    )
    return root, str(item["id"])


async def test_report_workspace_flow(tmp_path):
    root, template_id = _workspace(tmp_path)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        await client.post("/api/v1/workspaces/open", json={"path": str(root)})

        created = await client.post(
            "/api/v1/reports/drafts",
            json={
                "templateId": template_id,
                "templateScope": "project",
                "formIds": ["QZ-0001"],
                "header": {"reportTitle": "采购审计报告"},
            },
        )
        assert created.status_code == 201, created.text
        draft = created.json()
        report_id = draft["manifest"]["id"]
        assert draft["manifest"]["status"] == "draft"
        assert [s["id"] for s in draft["body"]["sections"]] == ["S1"]

        doc = (await client.get(f"/api/v1/reports/{report_id}/doc")).json()
        assert doc["skeleton"]["confidence"] == "high"
        assert doc["body"]["selectedFormIds"] == ["QZ-0001"]

        outline = await client.put(
            f"/api/v1/reports/{report_id}/outline",
            json={
                "approved": True,
                "sections": [
                    {"id": "S1", "title": "一、基本情况", "level": 1},
                    {"id": "S2", "title": "二、审计发现", "level": 1},
                ],
            },
        )
        assert outline.status_code == 200
        assert outline.json()["outlineApproved"] is True

        section = await client.put(
            f"/api/v1/reports/{report_id}/sections/S2",
            json={
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": "采购阿里云服务器，费用99元。",
                        "sourceRefs": ["QZ-0001"],
                    }
                ]
            },
        )
        assert section.status_code == 200
        assert section.json()["status"] == "written"

        # S1 还空着 → 校验失败、导出被拦截
        blocked = await client.post(f"/api/v1/reports/{report_id}/export", json={})
        assert blocked.status_code == 409
        assert blocked.json()["error"]["code"] == "E_VERIFY_FAILED"

        await client.put(
            f"/api/v1/reports/{report_id}/sections/S1",
            json={
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": "本次审计依据年度计划实施。",
                        "sourceRefs": ["QZ-0001"],
                    }
                ]
            },
        )
        check = (await client.post(f"/api/v1/reports/{report_id}/verify")).json()
        assert check["ok"] is True, check["errors"]

        exported = await client.post(f"/api/v1/reports/{report_id}/export", json={})
        assert exported.status_code == 200
        assert exported.json()["status"] == "saved"
        assert (
            root
            / ".fireeye"
            / "artifacts"
            / "reports"
            / report_id
            / "document.docx"
        ).is_file()

        items = (await client.get("/api/v1/reports")).json()["items"]
        assert any(item["id"] == report_id for item in items)

        messages = (await client.get(f"/api/v1/reports/{report_id}/messages")).json()
        assert messages["items"] == []

        await client.post("/api/v1/workspaces/close")


async def test_report_forms_endpoint(tmp_path):
    root, template_id = _workspace(tmp_path)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        await client.post("/api/v1/workspaces/open", json={"path": str(root)})
        created = await client.post(
            "/api/v1/reports/drafts",
            json={
                "templateId": template_id,
                "templateScope": "project",
                "formIds": ["QZ-0001"],
            },
        )
        report_id = created.json()["manifest"]["id"]

        cleared = await client.put(
            f"/api/v1/reports/{report_id}/forms", json={"formIds": []}
        )
        assert cleared.status_code == 200, cleared.text
        assert cleared.json()["selectedFormIds"] == []

        restored = await client.put(
            f"/api/v1/reports/{report_id}/forms", json={"formIds": ["QZ-0001"]}
        )
        assert restored.status_code == 200, restored.text
        assert restored.json()["selectedFormIds"] == ["QZ-0001"]

        doc = (await client.get(f"/api/v1/reports/{report_id}/doc")).json()
        assert doc["manifest"]["formIds"] == ["QZ-0001"]
        assert doc["manifest"]["formsSnapshot"][0]["id"] == "QZ-0001"

        await client.post("/api/v1/workspaces/close")


async def test_report_requires_workspace(tmp_path):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        await client.post("/api/v1/workspaces/close")
        response = await client.get("/api/v1/reports/RPT-0001/doc")
        assert response.status_code == 409
