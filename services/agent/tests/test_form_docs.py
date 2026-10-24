"""取证单草稿：创建、疑点范围、大纲、章节、标题与对话记录。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from fireeyeaud_agent.config import store as config_store
from fireeyeaud_agent.errors import ApiError
from fireeyeaud_agent.store import evidence_forms, form_docs, templates as templates_store
from fireeyeaud_agent.store.findings import create_finding


@pytest.fixture(autouse=True)
def _global_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))


def _template(root: Path) -> str:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("取证单")
    run.font.size = Pt(22)
    doc.add_paragraph("项目：{{项目名称}}")
    doc.add_paragraph("一、基本情况", style=doc.styles.add_style("标题 1", 1))
    doc.add_paragraph("二、取证内容", style=doc.styles.add_style("标题 2", 1))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.标题}}"

    item = config_store.create_item(
        "templates",
        {
            "scope": "project",
            "kind": "evidence-form",
            "name": "取证单模板",
            "fileName": "form.docx",
        },
        root,
    )
    target = templates_store.templates_root("project", root) / "evidence-form" / "form.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return str(item["id"])


def _finding(root: Path, title: str = "关联方交易") -> str:
    return create_finding(
        root,
        title=title,
        description="未披露",
        severity="high",
        dataRefs=[],
        sessionId="",
        origin="manual",
        evidenceRefs=["E-0001"],
    )["id"]


def _setup(tmp_path: Path) -> tuple[Path, str, str]:
    root = tmp_path / "ws"
    (root / ".fireeye").mkdir(parents=True)
    (root / ".fireeye" / "project.json").write_text(
        json.dumps({"name": "示例项目", "counters": {}}, ensure_ascii=False), "utf-8"
    )
    return root, _template(root), _finding(root)


def test_create_form_writes_workspace_files(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    created = form_docs.create_form(
        root,
        template_id=template_id,
        template_scope="project",
        finding_ids=[finding_id],
        header={"preparedBy": "张审计"},
        title="关联方取证单",
    )
    manifest, body, skeleton = created["manifest"], created["body"], created["skeleton"]
    assert manifest["id"] == "QZ-0001"
    assert manifest["kind"] == "evidence-form"
    assert manifest["status"] == "draft"
    assert manifest["title"] == "关联方取证单"
    assert manifest["findingIds"] == [finding_id]
    assert manifest["findingsSnapshot"][0]["title"] == "关联方交易"
    assert manifest["findingsSnapshot"][0]["evidence"] == ["E-0001"]
    assert manifest["findingCount"] == 1
    assert body["formId"] == "QZ-0001"
    assert body["selectedFindingIds"] == [finding_id]
    assert body["outlineApproved"] is False
    assert [s["title"] for s in body["sections"]] == [
        "一、基本情况",
        "二、取证内容",
    ]
    assert skeleton["templateText"]["blocks"]

    directory = evidence_forms.forms_dir(root) / "QZ-0001"
    for name in ("manifest.json", "body.json", "skeleton.json", "template.docx"):
        assert (directory / name).is_file()


def test_create_form_without_template_sections_uses_default_outline(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    # 覆盖模板为无标题样式
    doc = Document()
    doc.add_paragraph("纯文本模板。")
    target = templates_store.templates_root("project", root) / "evidence-form" / "form.docx"
    doc.save(str(target))
    created = form_docs.create_form(
        root, template_id=template_id, template_scope="project", finding_ids=[finding_id]
    )
    assert [s["title"] for s in created["body"]["sections"]] == [
        "一、基本情况",
        "二、取证内容",
        "三、结论与建议",
    ]


def test_create_form_rejects_unknown_finding_and_wrong_kind(tmp_path):
    root, template_id, _finding_id = _setup(tmp_path)
    with pytest.raises(ApiError) as exc:
        form_docs.create_form(
            root, template_id=template_id, template_scope="project", finding_ids=["F-9999"]
        )
    assert exc.value.code == "E_FINDINGS_REQUIRED"

    report_tpl = config_store.create_item(
        "templates",
        {"scope": "project", "kind": "report", "name": "报告模板", "fileName": "r.docx"},
        root,
    )
    (templates_store.templates_root("project", root) / "report").mkdir(
        parents=True, exist_ok=True
    )
    Document().save(
        str(templates_store.templates_root("project", root) / "report" / "r.docx")
    )
    with pytest.raises(ApiError) as exc2:
        form_docs.create_form(
            root,
            template_id=str(report_tpl["id"]),
            template_scope="project",
            finding_ids=[],
        )
    assert exc2.value.code == "E_TEMPLATE_KIND"


def test_set_findings_refreshes_scope_and_snapshot(tmp_path):
    root, template_id, first = _setup(tmp_path)
    second = _finding(root, "服务器采购")
    form_docs.create_form(
        root, template_id=template_id, template_scope="project", finding_ids=[first]
    )
    body = form_docs.set_findings(root, "QZ-0001", [first, second])
    assert body["selectedFindingIds"] == [first, second]
    doc = form_docs.get_document(root, "QZ-0001")
    assert doc["manifest"]["findingIds"] == [first, second]
    assert doc["manifest"]["findingCount"] == 2
    assert [item["title"] for item in doc["manifest"]["findingsSnapshot"]] == [
        "关联方交易",
        "服务器采购",
    ]


def test_outline_section_title_and_draft_marking(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    form_docs.create_form(
        root, template_id=template_id, template_scope="project", finding_ids=[finding_id]
    )
    body = form_docs.set_outline(
        root,
        "QZ-0001",
        [{"title": "一、事项经过", "level": 1}, {"title": "二、取证结论", "level": 1}],
        approved=True,
    )
    assert body["outlineApproved"] is True
    assert [s["id"] for s in body["sections"]] == ["S1", "S2"]

    section = form_docs.write_section(
        root,
        "QZ-0001",
        "S2",
        [{"type": "paragraph", "text": "关联方交易未披露。", "sourceRefs": [finding_id]}],
    )
    assert section["status"] == "written"

    with pytest.raises(ApiError) as exc:
        form_docs.write_section(root, "QZ-0001", "S9", [])
    assert exc.value.code == "E_SECTION_NOT_FOUND"

    body = form_docs.set_title(root, "QZ-0001", "新标题")
    assert body["title"] == "新标题"
    assert form_docs.get_document(root, "QZ-0001")["manifest"]["title"] == "新标题"


def test_chat_log_roundtrip(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    form_docs.create_form(
        root, template_id=template_id, template_scope="project", finding_ids=[finding_id]
    )
    form_docs.append_chat(root, "QZ-0001", {"role": "user", "content": "先给大纲"})
    form_docs.append_chat(root, "QZ-0001", {"role": "assistant", "content": "好的"})
    assert [e["content"] for e in form_docs.load_chat(root, "QZ-0001")] == [
        "先给大纲",
        "好的",
    ]