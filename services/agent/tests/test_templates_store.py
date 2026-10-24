"""模板存储：双作用域、默认解析、docx 体检与示例复制。"""

from pathlib import Path

import pytest
from docx import Document

from fireeyeaud_agent.errors import ApiError
from fireeyeaud_agent.store.templates import (
    default_template,
    delete_template,
    duplicate_template,
    ensure_samples,
    get_template,
    list_templates,
    read_template_content,
    replace_template_file,
    resolve_template,
    scan_docx,
    template_path,
    update_template,
    upload_template,
)


def _docx(path: Path, body: str = "项目名称：{{项目名称}}") -> Path:
    doc = Document()
    doc.add_paragraph(body)
    doc.add_paragraph("取证单编号：{{取证单编号}}")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(0, 2).text = "风险"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.编号}}"
    table.cell(1, 2).text = "{{疑点.标题}}"
    table.cell(2, 0).text = "{{/疑点}}"
    doc.save(path)
    return path


def _env(tmp_path, monkeypatch) -> Path:
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    root = tmp_path / "project"
    (root / ".fireeye").mkdir(parents=True, exist_ok=True)
    return root


def test_scan_docx(tmp_path):
    src = _docx(tmp_path / "t.docx")
    check = scan_docx(src, "evidence-form")
    assert "项目名称" in check["tags"]
    assert "疑点" in check["loops"]
    assert check["missing"] == []


def test_project_upload_list_default_delete(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)

    src = _docx(tmp_path / "标准取证单.docx")
    item = upload_template(root, src, kind="evidence-form", scope="project")
    assert item["id"] == "tpl_1"
    assert item["scope"] == "project"
    assert item["isDefault"] is True
    assert (
        root / ".fireeye" / "templates" / "evidence-form" / item["fileName"]
    ).is_file()

    # 元数据进入统一配置，不生成第二份清单
    assert (root / ".fireeye" / "config" / "templates.json").is_file()
    assert not (root / ".fireeye" / "templates" / "templates.json").exists()

    items = list_templates(root, "evidence-form")
    assert any(i["id"] == item["id"] and i["scope"] == "project" for i in items)

    delete_template(root, item["id"], "project")
    assert all(i["id"] != item["id"] for i in list_templates(root, "evidence-form"))


def test_upload_can_mark_default(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)

    first = upload_template(
        root, _docx(tmp_path / "第一份.docx"), kind="report", scope="project"
    )
    second = upload_template(
        root,
        _docx(tmp_path / "第二份.docx"),
        kind="report",
        scope="project",
        is_default=True,
    )
    assert first["isDefault"] is True
    assert second["isDefault"] is True

    items = [
        i
        for i in list_templates(root, "report")
        if i["scope"] == "project" and not i.get("builtin")
    ]
    assert [i["id"] for i in items if i["isDefault"]] == [second["id"]]
    assert default_template(root, "report")["id"] == second["id"]


def test_global_and_project_scope_and_default(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)

    g = upload_template(root, _docx(tmp_path / "全局.docx"), kind="report", scope="global")
    p = upload_template(root, _docx(tmp_path / "项目.docx"), kind="report", scope="project")

    items = list_templates(root, "report")
    assert {i["scope"] for i in items if not i.get("builtin")} == {"global", "project"}
    # 两个作用域各自可以有默认；解析时项目优先
    assert default_template(root, "report")["id"] == p["id"]
    delete_template(root, p["id"], "project")
    assert default_template(root, "report")["id"] == g["id"]


def test_samples_are_global_virtual_and_copyable(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)

    items = list_templates(root, "evidence-form")
    sample = next(i for i in items if i.get("builtin"))
    assert sample["scope"] == "global"
    assert template_path(root, sample["id"], "global").is_file()

    copied = duplicate_template(root, sample["id"], "global", target_scope="project")
    assert copied["scope"] == "project"
    assert copied["builtin"] is False
    assert template_path(root, copied["id"], "project").is_file()


# ---- 计划未覆盖、但实现必须成立的边界 ----

def test_scan_docx_reports_placeholders_without_warnings(tmp_path):
    """模板只作参考范本：未识别字段 / 缺少字段都不再算问题，占位符只做统计。"""
    doc = Document()
    doc.add_paragraph("{{项目名称}} {{某个任意字段}}")
    path = tmp_path / "odd.docx"
    doc.save(path)

    check = scan_docx(path, "evidence-form")
    assert check["tags"] == ["某个任意字段", "项目名称"]
    assert check["unknown"] == []
    assert check["missing"] == []
    assert check["ok"] is True


def test_scan_docx_never_blocks_on_unclosed_loops(tmp_path):
    """循环不再执行，未闭合的循环标记不影响生成。"""
    doc = Document()
    doc.add_paragraph("{{#疑点}}")
    doc.add_paragraph("{{疑点.标题}}")
    path = tmp_path / "loop.docx"
    doc.save(path)

    check = scan_docx(path, "evidence-form")
    assert check["loops"] == ["疑点"]
    assert check["unclosed"] == ["疑点"]
    assert check["balanced"] is False
    assert check["ok"] is True


def test_scan_docx_reads_headers_and_footers(tmp_path):
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "{{被审计单位}}"
    section.footer.paragraphs[0].text = "{{复核人}}"
    path = tmp_path / "hf.docx"
    doc.save(path)

    check = scan_docx(path, "report")
    assert "被审计单位" in check["tags"]
    assert "复核人" in check["tags"]


def test_upload_rejects_non_docx(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    bad = tmp_path / "x.txt"
    bad.write_text("hello", "utf-8")
    with pytest.raises(ApiError) as exc:
        upload_template(root, bad, kind="report", scope="project")
    assert exc.value.code == "E_FILE_TYPE"


def test_upload_rejects_missing_file(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    with pytest.raises(ApiError) as exc:
        upload_template(root, tmp_path / "nope.docx", kind="report", scope="project")
    assert exc.value.code == "E_FILE_NOT_FOUND"


def test_upload_rejects_oversized_file(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    big = tmp_path / "big.docx"
    big.write_bytes(b"0" * (20 * 1024 * 1024 + 1))
    with pytest.raises(ApiError) as exc:
        upload_template(root, big, kind="report", scope="project")
    assert exc.value.code == "E_FILE_TOO_LARGE"


def test_project_scope_requires_workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    with pytest.raises(ApiError) as exc:
        upload_template(None, _docx(tmp_path / "a.docx"), kind="report", scope="project")
    assert exc.value.code == "E_NO_WORKSPACE"


def test_project_item_overrides_same_global_id(tmp_path, monkeypatch):
    """load_effective 语义：同 id 的项目项覆盖全局项（P3 依赖这点）。"""
    root = _env(tmp_path, monkeypatch)
    g = upload_template(root, _docx(tmp_path / "g.docx"), kind="report", scope="global")

    from fireeyeaud_agent.config import store

    store.put_item(
        "templates",
        {
            **g,
            "id": g["id"],
            "scope": "project",
            "name": "项目覆盖",
            "fileName": g["fileName"],
        },
        root,
    )
    items = list_templates(root, "report")
    match = [i for i in items if i["id"] == g["id"]]
    assert len(match) == 1
    assert match[0]["scope"] == "project"
    assert match[0]["name"] == "项目覆盖"


def test_rename_does_not_change_scope_or_kind(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    item = upload_template(root, _docx(tmp_path / "a.docx"), kind="report", scope="project")
    updated = update_template(root, item["id"], "project", name="新名字")
    assert updated["name"] == "新名字"
    assert updated["scope"] == "project"
    assert updated["kind"] == "report"


def test_set_default_only_affects_same_scope(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    g1 = upload_template(root, _docx(tmp_path / "g1.docx"), kind="report", scope="global")
    g2 = upload_template(root, _docx(tmp_path / "g2.docx"), kind="report", scope="global")
    p1 = upload_template(root, _docx(tmp_path / "p1.docx"), kind="report", scope="project")

    update_template(root, g2["id"], "global", is_default=True)

    global_items = [i for i in list_templates(root, "report") if i["scope"] == "global"]
    assert [i["id"] for i in global_items if i["isDefault"]] == [g2["id"]]
    # 项目作用域的默认不受影响
    assert default_template(root, "report")["id"] == p1["id"]
    assert all(i["id"] != g1["id"] for i in global_items if i["isDefault"])


def test_delete_default_promotes_another_same_scope(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    first = upload_template(root, _docx(tmp_path / "1.docx"), kind="report", scope="project")
    second = upload_template(root, _docx(tmp_path / "2.docx"), kind="report", scope="project")
    assert default_template(root, "report")["id"] == first["id"]

    # second 设为默认后删除它 → first 重新成为默认
    update_template(root, second["id"], "project", is_default=True)
    assert default_template(root, "report")["id"] == second["id"]
    delete_template(root, second["id"], "project")
    assert default_template(root, "report")["id"] == first["id"]


def test_delete_builtin_is_refused(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    sample = next(i for i in list_templates(root, "report") if i.get("builtin"))
    with pytest.raises(ApiError) as exc:
        delete_template(root, sample["id"], "global")
    assert exc.value.code == "E_TEMPLATE_BUILTIN"


def test_default_template_falls_back_to_sample(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    resolved = default_template(root, "evidence-form")
    assert resolved is not None
    assert resolved["builtin"] is True


def test_duplicate_to_global_from_project(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    p = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")
    copied = duplicate_template(root, p["id"], "project", target_scope="global")
    assert copied["scope"] == "global"
    assert copied["id"] != p["id"]
    assert template_path(root, copied["id"], "global").is_file()


def test_resolve_and_get_template(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    p = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")

    assert get_template(root, p["id"], "global") is None
    assert get_template(root, p["id"], "project")["id"] == p["id"]
    # scope 缺省时按生效列表解析（P3 兼容）
    assert resolve_template(root, p["id"], None)["id"] == p["id"]
    with pytest.raises(ApiError):
        resolve_template(root, "nope", None)


def test_list_templates_fills_abs_path_for_global_and_project(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    g = upload_template(root, _docx(tmp_path / "g.docx"), kind="report", scope="global")
    p = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")

    by_id = {i["id"]: i for i in list_templates(root, "report")}
    assert Path(by_id[g["id"]]["absPath"]).is_file()
    assert Path(by_id[p["id"]]["absPath"]).is_file()
    # 全局模板位于项目目录之外
    assert ".fireeye" not in by_id[g["id"]]["absPath"]


def test_read_template_content_returns_docx_bytes(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    p = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")
    data, mime = read_template_content(root, p["id"], "project")
    assert data[:2] == b"PK"
    assert mime.startswith("application/vnd.openxmlformats")


def test_samples_are_not_written_to_any_manifest(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    ensure_samples(root)
    global_manifest = tmp_path / "global" / "config" / "templates.json"
    assert not global_manifest.exists()
    # 示例文件确实生成
    samples = tmp_path / "global" / "templates" / "samples"
    assert (samples / "evidence-form.docx").is_file()
    assert (samples / "report.docx").is_file()


def test_sample_documents_contain_expected_markers(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    ensure_samples(root)
    samples = tmp_path / "global" / "templates" / "samples"

    form = scan_docx(samples / "evidence-form.docx", "evidence-form")
    assert "项目名称" in form["tags"]
    assert "疑点" in form["loops"]
    assert form["missing"] == []
    assert form["unknown"] == []
    assert form["unclosed"] == []

    report = scan_docx(samples / "report.docx", "report")
    assert "报告编号" in report["tags"]
    assert "取证单" in report["loops"]
    assert report["missing"] == []
    assert report["unknown"] == []
    # 报告附录的顶层 `疑点` 循环必须有结束标记，否则 P3 渲染会报 Jinja 语法错
    assert report["unclosed"] == []
    assert report["balanced"] is True


def test_scan_docx_detects_unclosed_loop(tmp_path):
    doc = Document()
    doc.add_paragraph("{{项目名称}} {{报告编号}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "{{#疑点}}"
    table.cell(0, 1).text = "{{疑点.标题}}"
    path = tmp_path / "unclosed.docx"
    doc.save(path)

    check = scan_docx(path, "report")
    assert check["unclosed"] == ["疑点"]
    assert check["balanced"] is False


def test_scan_docx_reads_all_header_footer_variants(tmp_path):
    doc = Document()
    section = doc.sections[0]
    section.first_page_header.paragraphs[0].text = "{{被审计单位}}"
    section.even_page_footer.paragraphs[0].text = "{{复核人}}"
    section.first_page_footer.paragraphs[0].text = "{{编制人}}"
    path = tmp_path / "hf2.docx"
    doc.save(path)

    check = scan_docx(path, "report")
    assert {"被审计单位", "复核人", "编制人"} <= set(check["tags"])


def test_corrupt_sample_does_not_break_list(tmp_path, monkeypatch):
    """示例文件损坏/为空时，列表仍要可用（用户无法在 UI 里修复示例）。"""
    root = _env(tmp_path, monkeypatch)
    ensure_samples(root)
    broken = tmp_path / "global" / "templates" / "samples" / "report.docx"
    broken.write_bytes(b"")

    items = list_templates(root, "report")  # 不能抛异常
    assert isinstance(items, list)


def test_resolve_sample_with_explicit_global_scope(tmp_path, monkeypatch):
    """列表把示例标为 scope=global，显式作用域解析必须也能拿到它（P3 依赖）。"""
    root = _env(tmp_path, monkeypatch)
    sample = next(i for i in list_templates(root, "report") if i.get("builtin"))

    assert get_template(root, sample["id"], "global")["id"] == sample["id"]
    assert resolve_template(root, sample["id"], "global")["id"] == sample["id"]
    # 示例只存在于全局作用域
    assert get_template(root, sample["id"], "project") is None


def test_clearing_last_default_promotes_another(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    first = upload_template(root, _docx(tmp_path / "1.docx"), kind="report", scope="project")
    second = upload_template(root, _docx(tmp_path / "2.docx"), kind="report", scope="project")
    update_template(root, second["id"], "project", is_default=True)

    update_template(root, second["id"], "project", is_default=False)

    items = [
        i for i in list_templates(root, "report") if i["scope"] == "project" and not i.get("builtin")
    ]
    defaults = [i["id"] for i in items if i["isDefault"]]
    assert defaults == [first["id"]], "同一作用域/类型必须始终有一个默认"


def test_delete_global_default_keeps_project_default(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    g = upload_template(root, _docx(tmp_path / "g.docx"), kind="report", scope="global")
    p = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")
    assert default_template(root, "report")["id"] == p["id"]

    delete_template(root, g["id"], "global")

    assert default_template(root, "report")["id"] == p["id"]


def test_default_template_skips_missing_files(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    first = upload_template(root, _docx(tmp_path / "1.docx"), kind="report", scope="project")
    second = upload_template(root, _docx(tmp_path / "2.docx"), kind="report", scope="project")
    # 外部删掉默认项的文件
    template_path(root, first["id"], "project").unlink()

    assert default_template(root, "report")["id"] == second["id"]


def test_delete_template_rejects_unknown_scope(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    item = upload_template(root, _docx(tmp_path / "a.docx"), kind="report", scope="project")
    with pytest.raises(ApiError) as exc:
        delete_template(root, item["id"], "bogus")
    assert exc.value.code == "E_INVALID_SCOPE"


# ---- 修改作用域 = 搬家 ----


def test_move_project_to_global(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    item = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")
    old_path = template_path(root, item["id"], "project")
    assert old_path.is_file()

    moved = update_template(root, item["id"], "project", target_scope="global")

    assert moved["scope"] == "global"
    assert not old_path.exists(), "搬家必须删除原作用域的文件"
    assert template_path(root, item["id"], "global").is_file()
    assert all(
        i["id"] != item["id"]
        for i in list_templates(root, "report")
        if i["scope"] == "project"
    )
    # 搬家后全局可见
    assert any(
        i["id"] == item["id"] for i in list_templates(root, "report") if i["scope"] == "global"
    )


def test_move_global_to_project_is_project_only(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    item = upload_template(root, _docx(tmp_path / "g.docx"), kind="report", scope="global")

    update_template(root, item["id"], "global", target_scope="project")

    items = {i["id"]: i for i in list_templates(root, "report")}
    assert items[item["id"]]["scope"] == "project"
    assert template_path(root, item["id"], "global") is None
    assert template_path(root, item["id"], "project").is_file()


def test_move_transfers_default_and_promotes_source(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    first = upload_template(root, _docx(tmp_path / "1.docx"), kind="report", scope="project")
    second = upload_template(root, _docx(tmp_path / "2.docx"), kind="report", scope="project")
    update_template(root, second["id"], "project", is_default=True)

    moved = update_template(root, second["id"], "project", target_scope="global")

    assert moved["isDefault"] is True, "默认标记随模板一起搬走"
    # 原作用域仍有健康默认
    assert default_template(root, "report")["id"] == first["id"]


def test_move_without_workspace_to_project_is_rejected(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    item = upload_template(None, _docx(tmp_path / "g.docx"), kind="report", scope="global")
    with pytest.raises(ApiError) as exc:
        update_template(None, item["id"], "global", target_scope="project")
    assert exc.value.code == "E_NO_WORKSPACE"


def test_move_rejects_missing_file(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    item = upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")
    template_path(root, item["id"], "project").unlink()

    with pytest.raises(ApiError) as exc:
        update_template(root, item["id"], "project", target_scope="global")
    assert exc.value.code == "E_FILE_NOT_FOUND"


# ---- 重新上传覆盖原文件 ----


def test_replace_template_file_keeps_identity_and_updates_check(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    item = upload_template(root, _docx(tmp_path / "old.docx"), kind="report", scope="project")
    old_path = template_path(root, item["id"], "project")

    new_src = Document()
    new_src.add_paragraph("{{项目名称}}")
    new_src.add_paragraph("{{报告编号}}")
    new_src.add_paragraph("{{取证单数量}}")
    new_path = tmp_path / "new.docx"
    new_src.save(new_path)

    replaced = replace_template_file(root, item["id"], "project", new_path)

    assert replaced["id"] == item["id"], "覆盖文件不换 id"
    assert replaced["name"] == item["name"], "覆盖文件不改名称"
    assert replaced["scope"] == "project"
    assert "取证单数量" in replaced["check"]["tags"]
    assert template_path(root, item["id"], "project") == old_path
    assert old_path.is_file()


def test_replace_rejects_builtin_and_bad_files(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    sample = next(i for i in list_templates(root, "report") if i.get("builtin"))
    good = _docx(tmp_path / "g.docx")
    with pytest.raises(ApiError) as exc:
        replace_template_file(root, sample["id"], "global", good)
    assert exc.value.code == "E_TEMPLATE_BUILTIN"

    item = upload_template(root, _docx(tmp_path / "a.docx"), kind="report", scope="project")
    bad = tmp_path / "x.txt"
    bad.write_text("hi", "utf-8")
    with pytest.raises(ApiError) as exc2:
        replace_template_file(root, item["id"], "project", bad)
    assert exc2.value.code == "E_FILE_TYPE"


def test_template_path_rejects_path_traversal(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    assert template_path(root, "../../etc/passwd", "project") is None
    assert template_path(root, "tpl_1", "project") is None  # 尚未上传


def test_list_templates_sorted_defaults_first(tmp_path, monkeypatch):
    root = _env(tmp_path, monkeypatch)
    first = upload_template(root, _docx(tmp_path / "1.docx"), kind="report", scope="project")
    upload_template(root, _docx(tmp_path / "2.docx"), kind="report", scope="project")

    items = [
        i for i in list_templates(root, "report") if not i.get("builtin")
    ]
    assert items[0]["id"] == first["id"]
    assert items[0]["isDefault"] is True


def test_same_file_name_uploads_do_not_overwrite(tmp_path, monkeypatch):
    """文件名用模板 id，两份同名 docx 都保留。"""
    root = _env(tmp_path, monkeypatch)
    d1 = tmp_path / "a"
    d2 = tmp_path / "b"
    d1.mkdir()
    d2.mkdir()
    one = upload_template(root, _docx(d1 / "同名.docx"), kind="report", scope="project")
    two = upload_template(root, _docx(d2 / "同名.docx"), kind="report", scope="project")

    assert one["id"] != two["id"]
    assert one["fileName"] != two["fileName"]
    assert template_path(root, one["id"], "project").is_file()
    assert template_path(root, two["id"], "project").is_file()


def test_no_template_counter_in_project_json(tmp_path, monkeypatch):
    """模板计数走全局 `_counters.json`，不污染项目 project.json。"""
    root = _env(tmp_path, monkeypatch)
    import json

    project = root / ".fireeye" / "project.json"
    project.parent.mkdir(parents=True, exist_ok=True)
    project.write_text(json.dumps({"counters": {"finding": 0}}), "utf-8")

    upload_template(root, _docx(tmp_path / "p.docx"), kind="report", scope="project")

    data = json.loads(project.read_text("utf-8"))
    assert "template" not in json.dumps(data.get("counters", {}))
    counters = tmp_path / "global" / "config" / "_counters.json"
    assert "templates" in counters.read_text("utf-8")