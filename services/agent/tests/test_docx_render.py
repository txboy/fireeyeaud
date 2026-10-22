"""模板语法 → docxtpl 标记的预处理与渲染。"""

from pathlib import Path

from docx import Document
from docx.shared import Inches

from fireeyeaud_agent.store.docx_render import prepare_template, render_docx


def _template(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("项目：{{项目名称}}")
    doc.add_paragraph("取证单：{{取证单编号}}")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.编号}} {{疑点.标题}}"
    table.cell(2, 0).text = "{{/疑点}}"
    doc.save(path)
    return path


def test_prepare_and_render(tmp_path):
    src = _template(tmp_path / "tpl.docx")
    prepared = tmp_path / "prepared.docx"
    out = tmp_path / "out.docx"
    prepare_template(src, prepared, "evidence-form")

    render_docx(
        prepared,
        out,
        {
            "项目名称": "年报审计",
            "取证单编号": "QZ-0001",
            "findings": [
                {"编号": "F-0001", "标题": "关联方交易"},
                {"编号": "F-0002", "标题": "收入截止"},
            ],
        },
    )

    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "年报审计" in text
    table_text = "\n".join(
        cell.text
        for table in Document(out).tables
        for row in table.rows
        for cell in row.cells
    )
    assert "F-0001" in table_text
    assert "F-0002" in table_text
    assert "关联方交易" in table_text


def test_prepare_maps_paragraph_loop(tmp_path):
    """段落循环（非表格）应变成 {%p for ... %}，并保留段落。"""
    doc = Document()
    doc.add_paragraph("项目：{{项目名称}}")
    doc.add_paragraph("{{#疑点}}")
    doc.add_paragraph("{{疑点.标题}}")
    doc.add_paragraph("{{/疑点}}")
    src = tmp_path / "p.docx"
    doc.save(src)

    prepared = tmp_path / "prepared.docx"
    prepare_template(src, prepared, "evidence-form")
    rendered = tmp_path / "out.docx"
    render_docx(
        prepared,
        rendered,
        {"项目名称": "P", "findings": [{"标题": "A"}, {"标题": "B"}]},
    )

    lines = [p.text for p in Document(rendered).paragraphs]
    assert "A" in lines
    assert "B" in lines
    assert not any("{{" in line for line in lines)


def test_prepare_handles_form_loop_for_reports(tmp_path):
    doc = Document()
    doc.add_paragraph("报告：{{报告编号}}")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#取证单}}"
    table.cell(1, 1).text = "{{取证单.编号}} {{取证单.疑点数量}}"
    table.cell(2, 0).text = "{{/取证单}}"
    src = tmp_path / "r.docx"
    doc.save(src)

    prepared = tmp_path / "prepared.docx"
    prepare_template(src, prepared, "report")
    out = tmp_path / "out.docx"
    render_docx(
        prepared,
        out,
        {
            "报告编号": "RPT-0001",
            "forms": [{"编号": "QZ-0001", "疑点数量": 2}, {"编号": "QZ-0002", "疑点数量": 1}],
        },
    )

    cells = [c.text for t in Document(out).tables for r in t.rows for c in r.cells]
    joined = "\n".join(cells)
    assert "QZ-0001" in joined
    assert "QZ-0002" in joined
    # 标记单元格在循环体里被清空，数据落在后面的单元格
    assert all("{{" not in c and "{%" not in c for c in cells)


def test_prepare_rejects_unsupported_loop(tmp_path):
    import pytest

    from fireeyeaud_agent.errors import ApiError

    doc = Document()
    doc.add_paragraph("{{#未知循环}}")
    doc.add_paragraph("{{/未知循环}}")
    src = tmp_path / "bad.docx"
    doc.save(src)

    with pytest.raises(ApiError) as exc:
        prepare_template(src, tmp_path / "out.docx", "report")
    assert exc.value.code == "E_TEMPLATE_LOOP"


def test_render_escapes_xml_special_characters(tmp_path):
    """`&` 与 `<` 必须被转义：否则 docx 的 XML 会被破坏，内容静默丢失。"""
    doc = Document()
    doc.add_paragraph("单位：{{项目名称}}")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "{{#疑点}}"
    table.cell(1, 0).text = "{{疑点.标题}}"
    src = tmp_path / "amp.docx"
    doc.save(src)

    # 段落循环：把 table 换成段落，便于断言
    doc2 = Document()
    doc2.add_paragraph("单位：{{项目名称}}")
    doc2.add_paragraph("{{#疑点}}")
    doc2.add_paragraph("{{疑点.标题}}")
    doc2.add_paragraph("{{/疑点}}")
    src2 = tmp_path / "amp2.docx"
    doc2.save(src2)

    prepared = tmp_path / "prepared.docx"
    prepare_template(src2, prepared, "evidence-form")
    out = tmp_path / "out.docx"
    render_docx(
        prepared,
        out,
        {
            "项目名称": "A&B <100万",
            "findings": [{"标题": "P&L <阈值>"}],
        },
    )

    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "A&B <100万" in text
    assert "P&L <阈值>" in text


def test_render_rejects_sandboxed_jinja_escapes(tmp_path):
    """用户上传的模板不可执行任意 Python（沙箱环境）。"""
    import pytest

    from fireeyeaud_agent.errors import ApiError

    doc = Document()
    doc.add_paragraph("{{ ''.__class__.__mro__ }}")
    src = tmp_path / "evil.docx"
    doc.save(src)
    prepared = tmp_path / "prepared.docx"
    prepare_template(src, prepared, "report")

    with pytest.raises(ApiError) as exc:
        render_docx(prepared, tmp_path / "out.docx", {})
    assert exc.value.code == "E_TEMPLATE_RENDER"


def test_header_table_loop_repeats_rows(tmp_path):
    """页眉表格里的行循环也必须克隆数据行，否则只会剩下表头。"""
    doc = Document()
    section = doc.sections[0]
    table = section.header.add_table(rows=3, cols=2, width=Inches(6))
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.标题}}"
    table.cell(2, 0).text = "{{/疑点}}"
    src = tmp_path / "hdr.docx"
    doc.save(src)

    prepared = tmp_path / "prepared.docx"
    prepare_template(src, prepared, "evidence-form")
    out = tmp_path / "out.docx"
    render_docx(
        prepared,
        out,
        {"findings": [{"标题": "甲"}, {"标题": "乙"}]},
    )

    header_tables = Document(out).sections[0].header.tables
    cells = [c.text for t in header_tables for r in t.rows for c in r.cells]
    assert "甲" in cells
    assert "乙" in cells