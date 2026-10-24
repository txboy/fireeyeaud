"""取证单导出：模板样式继承、占位符不残留、导出后状态与过期标记。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from fireeyeaud_agent.config import store as config_store
from fireeyeaud_agent.errors import ApiError
from fireeyeaud_agent.store import evidence_forms, form_docs, templates as templates_store
from fireeyeaud_agent.store.findings import create_finding


@pytest.fixture(autouse=True)
def _global_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))


def _east_asia(style_or_run, name: str) -> None:
    element = (
        style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    )
    element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _setup(tmp_path: Path) -> tuple[Path, str, str]:
    root = tmp_path / "ws"
    (root / ".fireeye").mkdir(parents=True)
    (root / ".fireeye" / "project.json").write_text(
        json.dumps({"name": "示例项目", "counters": {}}, ensure_ascii=False), "utf-8"
    )
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(16)
    _east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = Pt(30)
    heading = doc.styles.add_style("标题 1", 1)
    heading.base_style = normal
    heading.font.bold = True
    _east_asia(heading, "黑体")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("取证单范本")
    run.font.size = Pt(22)
    doc.add_paragraph("项目：{{项目名称}}，编号 {{取证单编号}}")
    doc.add_paragraph("一、基本情况", style=heading)
    doc.add_paragraph("此处填写基本情况。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.标题}}"

    item = config_store.create_item(
        "templates",
        {"scope": "project", "kind": "evidence-form", "name": "模板", "fileName": "f.docx"},
        root,
    )
    target = templates_store.templates_root("project", root) / "evidence-form" / "f.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))

    finding_id = create_finding(
        root,
        title="关联方交易",
        description="未披露，金额99元。",
        severity="high",
        dataRefs=[],
        sessionId="",
        origin="manual",
    )["id"]
    return root, str(item["id"]), finding_id


def _prepare(root: Path, template_id: str, finding_id: str) -> None:
    form_docs.create_form(
        root, template_id=template_id, template_scope="project", finding_ids=[finding_id]
    )
    form_docs.set_outline(
        root, "QZ-0001", [{"title": "一、事项经过", "level": 1}], approved=True
    )
    form_docs.write_section(
        root,
        "QZ-0001",
        "S1",
        [
            {
                "type": "paragraph",
                "text": "关联方交易未披露，金额99元。",
                "sourceRefs": [finding_id],
            }
        ],
    )


def test_export_blocks_when_verification_fails(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    form_docs.create_form(
        root, template_id=template_id, template_scope="project", finding_ids=[finding_id]
    )
    with pytest.raises(ApiError) as exc:
        form_docs.export_form(root, "QZ-0001")
    assert exc.value.code == "E_VERIFY_FAILED"
    assert not (evidence_forms.forms_dir(root) / "QZ-0001" / "document.docx").exists()


def test_export_inherits_template_styles_and_drops_placeholders(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    _prepare(root, template_id, finding_id)

    manifest = form_docs.export_form(root, "QZ-0001")
    assert manifest["status"] == "saved"
    assert "document.docx" in manifest["files"]

    exported = Document(
        str(evidence_forms.forms_dir(root) / "QZ-0001" / "document.docx")
    )
    filled = [p for p in exported.paragraphs if p.text.strip()]
    assert filled[0].text == "示例项目取证单"
    assert filled[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert filled[0].runs[0].font.size.pt == 22

    heading = next(p for p in filled if p.text == "一、事项经过")
    assert heading.runs[0].font.bold is True
    assert heading.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"

    body = next(p for p in filled if "关联方交易未披露" in p.text)
    assert body.runs[0].font.size.pt == 16
    assert body.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "仿宋_GB2312"

    text = "\n".join(p.text for p in filled)
    assert "{{" not in text
    assert "此处填写" not in text
    # 占位符表格不得残留
    assert all(
        "{{疑点" not in cell.text
        for table in exported.tables
        for row in table.rows
        for cell in row.cells
    )


def test_editing_after_export_marks_draft_again(tmp_path):
    root, template_id, finding_id = _setup(tmp_path)
    _prepare(root, template_id, finding_id)
    form_docs.export_form(root, "QZ-0001")
    assert form_docs.get_document(root, "QZ-0001")["manifest"]["status"] == "saved"

    form_docs.write_section(
        root,
        "QZ-0001",
        "S1",
        [
            {
                "type": "paragraph",
                "text": "关联方交易未披露，金额99元。",
                "sourceRefs": [finding_id],
                "origin": "user",
            }
        ],
    )
    assert form_docs.get_document(root, "QZ-0001")["manifest"]["status"] == "draft"