"""取证单：预览渲染、保存落盘、列表与删除。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from fireeyeaud_agent.errors import ApiError
from fireeyeaud_agent.store.evidence_forms import (
    attach_receipt,
    clear_receipt,
    commit_evidence_form,
    delete_evidence_form,
    get_evidence_form,
    list_evidence_forms,
    preview_evidence_form,
)
from fireeyeaud_agent.store.findings import create_finding
from fireeyeaud_agent.store.templates import upload_template


def _template(path: Path, kind: str = "evidence-form") -> Path:
    doc = Document()
    doc.add_paragraph("项目：{{项目名称}}")
    doc.add_paragraph("编号：{{取证单编号}}")
    doc.add_paragraph("标题：{{取证单标题}}")
    doc.add_paragraph("数量：{{疑点数量}}")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "标题"
    table.cell(1, 0).text = "{{#疑点}}"
    table.cell(1, 1).text = "{{疑点.编号}} {{疑点.标题}}"
    table.cell(2, 0).text = "{{/疑点}}"
    doc.save(path)
    return path


def _setup(tmp_path: Path, monkeypatch) -> tuple[str, str]:
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    (tmp_path / ".fireeye").mkdir(parents=True, exist_ok=True)
    fid = create_finding(
        tmp_path,
        title="关联方交易",
        description="未披露",
        severity="high",
        dataRefs=[],
        sessionId="",
        origin="manual",
    )["id"]
    tpl = upload_template(
        tmp_path, _template(tmp_path / "tpl.docx"), kind="evidence-form", scope="project"
    )
    return fid, tpl["id"]


def test_preview_and_commit(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={"projectName": "年报审计", "preparedBy": "张审计"},
    )
    assert (tmp_path / preview["relPath"]).is_file()
    assert preview["suggestedTitle"]
    assert preview["kind"] == "evidence-form"

    manifest = commit_evidence_form(tmp_path, preview["token"], title="关联方取证单")
    assert manifest["id"] == "QZ-0001"
    assert manifest["templateScope"] == "project"
    assert manifest["templateId"] == tpl_id
    assert manifest["findingIds"] == [fid]
    assert manifest["findingsSnapshot"][0]["title"] == "关联方交易"
    artifact = tmp_path / ".fireeye" / "artifacts" / "evidence-forms" / "QZ-0001"
    assert (artifact / "document.docx").is_file()
    assert (artifact / "template.docx").is_file()
    assert (artifact / "manifest.json").is_file()
    assert list_evidence_forms(tmp_path)[0]["id"] == "QZ-0001"

    # 保存后暂存区删除
    assert not (tmp_path / ".fireeye" / "staging" / preview["token"]).exists()


def test_preview_writes_rendered_content(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={"projectName": "年报审计"},
    )
    rendered = Document(str(tmp_path / preview["relPath"]))
    text = "\n".join(p.text for p in rendered.paragraphs)
    assert "年报审计" in text
    table_text = "\n".join(
        cell.text for table in rendered.tables for row in table.rows for cell in row.cells
    )
    assert "关联方交易" in table_text


def test_preview_rejects_bad_inputs(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)

    with pytest.raises(ApiError) as exc:
        preview_evidence_form(
            tmp_path,
            template_id=tpl_id,
            template_scope="project",
            finding_ids=[],
            header={},
        )
    assert exc.value.code == "E_FINDINGS_REQUIRED"

    with pytest.raises(ApiError) as exc2:
        preview_evidence_form(
            tmp_path,
            template_id=tpl_id,
            template_scope="project",
            finding_ids=["F-9999"],
            header={},
        )
    assert exc2.value.code == "E_FINDINGS_REQUIRED"

    # 显式作用域找不到模板时不跨作用域回退
    with pytest.raises(ApiError) as exc3:
        preview_evidence_form(
            tmp_path,
            template_id=tpl_id,
            template_scope="global",
            finding_ids=[fid],
            header={},
        )
    assert exc3.value.code == "E_NOT_FOUND"


def test_preview_rejects_wrong_template_kind(tmp_path, monkeypatch):
    fid, _ = _setup(tmp_path, monkeypatch)
    report = upload_template(
        tmp_path, _template(tmp_path / "r.docx"), kind="report", scope="project"
    )
    with pytest.raises(ApiError) as exc:
        preview_evidence_form(
            tmp_path,
            template_id=report["id"],
            template_scope="project",
            finding_ids=[fid],
            header={},
        )
    assert exc.value.code == "E_TEMPLATE_KIND"


def test_warnings_include_template_check_and_missing_evidence(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))
    (tmp_path / ".fireeye").mkdir(parents=True, exist_ok=True)
    fid = create_finding(
        tmp_path,
        title="无证据疑点",
        description="",
        severity="low",
        dataRefs=[],
        sessionId="",
        origin="manual",
    )["id"]
    # 缺少 {{取证单编号}}，且有未识别字段
    doc = Document()
    doc.add_paragraph("{{项目名称}} {{未知字段}}")
    doc.save(tmp_path / "warn.docx")
    tpl = upload_template(
        tmp_path, tmp_path / "warn.docx", kind="evidence-form", scope="project"
    )

    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl["id"],
        template_scope="project",
        finding_ids=[fid],
        header={"projectName": "P"},
    )
    joined = "\n".join(preview["warnings"])
    assert "取证单编号" in joined
    assert "未知字段" in joined


def test_commit_expired_staging(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={},
    )
    payload_path = tmp_path / ".fireeye" / "staging" / preview["token"] / "payload.json"
    payload = json.loads(payload_path.read_text("utf-8"))
    payload["createdAt"] = "2000-01-01T00:00:00+00:00"
    payload_path.write_text(json.dumps(payload), "utf-8")

    with pytest.raises(ApiError) as exc:
        commit_evidence_form(tmp_path, preview["token"])
    assert exc.value.code == "E_STAGING_EXPIRED"


def test_unknown_token_and_delete(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    with pytest.raises(ApiError):
        commit_evidence_form(tmp_path, "nope")

    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={},
    )
    manifest = commit_evidence_form(tmp_path, preview["token"])
    assert get_evidence_form(tmp_path, manifest["id"])["id"] == manifest["id"]

    delete_evidence_form(tmp_path, manifest["id"])
    assert list_evidence_forms(tmp_path) == []
    assert not (
        tmp_path / ".fireeye" / "artifacts" / "evidence-forms" / manifest["id"]
    ).exists()
    with pytest.raises(ApiError):
        get_evidence_form(tmp_path, manifest["id"])


def test_receipt_upload_replace_delete(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={"projectName": "P"},
    )
    manifest = commit_evidence_form(tmp_path, preview["token"], title="取证单")
    form_id = manifest["id"]

    a = tmp_path / "客户签收.pdf"
    a.write_bytes(b"%PDF a")
    updated = attach_receipt(tmp_path, form_id, a)
    assert updated["receipt"]["name"] == "客户签收.pdf"
    assert updated["receipt"]["path"] == f"回执单/{form_id}_客户签收.pdf"
    assert (tmp_path / "回执单" / f"{form_id}_客户签收.pdf").is_file()

    b = tmp_path / "客户签收2.pdf"
    b.write_bytes(b"%PDF b")
    updated = attach_receipt(tmp_path, form_id, b)
    assert updated["receipt"]["name"] == "客户签收2.pdf"
    assert not (tmp_path / "回执单" / f"{form_id}_客户签收.pdf").exists()

    assert clear_receipt(tmp_path, form_id)["receipt"] is None
    assert not (tmp_path / "回执单" / f"{form_id}_客户签收2.pdf").exists()


def test_receipt_rejects_bad_type_and_missing_form(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={},
    )
    manifest = commit_evidence_form(tmp_path, preview["token"])

    bad = tmp_path / "x.exe"
    bad.write_bytes(b"x")
    with pytest.raises(ApiError) as exc:
        attach_receipt(tmp_path, manifest["id"], bad)
    assert exc.value.code == "E_FILE_TYPE"

    with pytest.raises(ApiError):
        attach_receipt(tmp_path, "QZ-9999", bad)


def test_delete_form_removes_receipt(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={},
    )
    manifest = commit_evidence_form(tmp_path, preview["token"])
    receipt = tmp_path / "r.pdf"
    receipt.write_bytes(b"%PDF")
    attach_receipt(tmp_path, manifest["id"], receipt)

    delete_evidence_form(tmp_path, manifest["id"])
    assert not (tmp_path / "回执单").exists() or not any(
        (tmp_path / "回执单").iterdir()
    )


def test_artifact_keeps_template_snapshot_after_template_deleted(tmp_path, monkeypatch):
    fid, tpl_id = _setup(tmp_path, monkeypatch)
    preview = preview_evidence_form(
        tmp_path,
        template_id=tpl_id,
        template_scope="project",
        finding_ids=[fid],
        header={},
    )
    manifest = commit_evidence_form(tmp_path, preview["token"])

    from fireeyeaud_agent.store.templates import delete_template

    delete_template(tmp_path, tpl_id, "project")
    artifact = tmp_path / manifest["artifactRef"]
    assert (artifact / "template.docx").is_file()
    assert (artifact / "document.docx").is_file()