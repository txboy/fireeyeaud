"""报告草稿：创建、大纲、章节、导出与校验。"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from fireeyeaud_agent.config import store as config_store
from fireeyeaud_agent.errors import ApiError
from fireeyeaud_agent.store import evidence_forms
from fireeyeaud_agent.store import report_docs, templates as templates_store
from fireeyeaud_agent.store.report_verify import verify_report


@pytest.fixture(autouse=True)
def _global_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("FIREYEAUD_GLOBAL_DIR", str(tmp_path / "global"))


def _east_asia(style_or_run, name: str) -> None:
    element = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _template(path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.8)
    section.top_margin = Cm(3.7)

    normal = doc.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(16)
    _east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = Pt(30)
    normal.paragraph_format.first_line_indent = Pt(32)

    heading = doc.styles.add_style("标题 1", 1)
    heading.base_style = normal
    heading.font.name = "黑体"
    heading.font.size = Pt(16)
    heading.font.bold = True
    _east_asia(heading, "黑体")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("审计报告模板标题")
    run.font.name = "小标宋"
    run.font.size = Pt(22)
    _east_asia(run, "小标宋")

    doc.add_paragraph("一、基本情况", style=heading)
    doc.add_paragraph("此处填写基本情况。")
    doc.add_paragraph("二、审计发现", style=heading)
    doc.add_paragraph("此处填写审计发现。")

    table = doc.add_table(rows=1, cols=3)
    table.style = doc.styles["Table Grid"]
    for index, text in enumerate(["序号", "疑点", "金额"]):
        table.rows[0].cells[index].text = text
        table.rows[0].cells[index].paragraphs[0].runs[0].font.bold = True

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _workspace(tmp_path: Path) -> tuple[Path, str]:
    root = tmp_path / "ws"
    (root / ".fireeye").mkdir(parents=True)
    (root / ".fireeye" / "project.json").write_text(
        json.dumps({"name": "示例项目", "counters": {}}, ensure_ascii=False), "utf-8"
    )
    item = config_store.create_item(
        "templates",
        {
            "scope": "project",
            "kind": "report",
            "name": "测试报告模板",
            "fileName": "tpl_report.docx",
        },
        root,
    )
    target = (
        templates_store.templates_root("project", root)
        / "report"
        / "tpl_report.docx"
    )
    _template(target)
    return root, str(item["id"])


def _form(root: Path, form_id: str = "QZ-0001", amount: str = "99") -> str:
    directory = evidence_forms.forms_dir(root) / form_id
    directory.mkdir(parents=True, exist_ok=True)
    manifest = {
        "id": form_id,
        "title": "阿里云服务器采购取证单",
        "createdAt": "2026-10-22T10:00:00+08:00",
        "header": {"preparedBy": "张审计"},
        "receipt": {"name": "回执.pdf", "path": "回执单/a.pdf", "uploadedAt": "2026-10-22"},
        "findingsSnapshot": [
            {
                "id": "F-0001",
                "title": "服务器采购费用",
                "description": f"采购阿里云服务器，费用{amount}元。",
                "severity": "low",
                "status": "confirmed",
                "subjects": ["信息化费"],
            }
        ],
    }
    (directory / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False), "utf-8"
    )
    return form_id


def _create(root: Path, template_id: str, form_ids: list[str]) -> dict:
    return report_docs.create_report(
        root,
        template_id=template_id,
        template_scope="project",
        form_ids=form_ids,
        header={"reportTitle": "阿里云服务器采购审计报告"},
    )


def test_create_report_uses_template_outline(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    created = _create(root, template_id, [form_id])

    manifest = created["manifest"]
    body = created["body"]
    assert manifest["id"] == "RPT-0001"
    assert manifest["status"] == "draft"
    assert manifest["title"] == "阿里云服务器采购审计报告"
    assert [s["title"] for s in body["sections"]] == ["一、基本情况", "二、审计发现"]
    assert body["outlineApproved"] is False
    assert body["selectedFormIds"] == [form_id]
    assert created["skeleton"]["confidence"] == "high"

    directory = report_docs._dir(root, "RPT-0001")
    assert (directory / "template.docx").is_file()
    assert (directory / "body.json").is_file()


def test_outline_and_section_write(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])

    body = report_docs.set_outline(
        root,
        "RPT-0001",
        [
            {"title": "一、项目概况", "level": 1, "goal": "写清背景"},
            {"title": "二、发现的问题", "level": 1},
            {"title": "三、结论", "level": 1},
        ],
        approved=True,
    )
    assert body["outlineApproved"] is True
    assert [s["id"] for s in body["sections"]] == ["S1", "S2", "S3"]

    section = report_docs.write_section(
        root,
        "RPT-0001",
        "S2",
        [
            {
                "type": "paragraph",
                "text": "采购阿里云服务器，费用99元。",
                "sourceRefs": [form_id],
            },
            {"type": "bullets", "items": ["费用已在信息化费列支"], "sourceRefs": [form_id]},
            {
                "type": "table",
                "columns": ["序号", "金额"],
                "rows": [["1", "99元"]],
                "sourceRefs": [form_id],
            },
        ],
    )
    assert section["status"] == "written"
    assert len(section["blocks"]) == 3

    with pytest.raises(ApiError) as missing:
        report_docs.write_section(root, "RPT-0001", "S9", [])
    assert missing.value.code == "E_SECTION_NOT_FOUND"


def test_outline_preserves_blocks_for_same_id(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    report_docs.write_section(
        root, "RPT-0001", "S1", [{"type": "paragraph", "text": "第一版", "sourceRefs": [form_id]}]
    )
    body = report_docs.set_outline(
        root,
        "RPT-0001",
        [
            {"id": "S1", "title": "一、项目概况", "level": 1},
            {"id": "S2", "title": "二、发现", "level": 1},
        ],
        approved=True,
    )
    assert body["sections"][0]["title"] == "一、项目概况"
    assert body["sections"][0]["blocks"][0]["text"] == "第一版"


def _fill_all(root: Path, form_id: str) -> None:
    body = report_docs.set_outline(
        root,
        "RPT-0001",
        [
            {"id": "S1", "title": "一、基本情况", "level": 1},
            {"id": "S2", "title": "二、审计发现", "level": 1},
        ],
        approved=True,
    )
    for section in body["sections"]:
        report_docs.write_section(
            root,
            "RPT-0001",
            section["id"],
            [
                {
                    "type": "paragraph",
                    "text": "采购阿里云服务器，费用99元，已在信息化费列支。",
                    "sourceRefs": [form_id],
                }
            ],
        )


def test_export_writes_document_with_template_styles(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    _fill_all(root, form_id)

    check = verify_report(root, "RPT-0001")
    assert check["ok"], check["errors"]

    manifest = report_docs.export_report(root, "RPT-0001")
    assert manifest["status"] == "saved"
    assert "document.docx" in manifest["files"]

    path = report_docs._dir(root, "RPT-0001") / "document.docx"
    exported = Document(str(path))
    filled = [p for p in exported.paragraphs if p.text.strip()]
    assert filled[0].text == "阿里云服务器采购审计报告"
    assert filled[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert filled[0].runs[0].font.size.pt == 22
    assert filled[0].runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "小标宋"

    heading = next(p for p in filled if p.text == "一、基本情况")
    assert heading.style.name == "标题 1"
    assert heading.runs[0].font.size.pt == 16
    assert heading.runs[0].font.bold is True
    assert heading.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"

    body_para = next(p for p in filled if "采购阿里云服务器" in p.text)
    assert body_para.runs[0].font.size.pt == 16
    assert body_para.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "仿宋_GB2312"
    assert round(body_para.paragraph_format.first_line_indent.pt) == 32
    assert round(body_para.paragraph_format.line_spacing.pt) == 30

    # 模板页边距保留
    assert round(exported.sections[0].left_margin.cm, 1) == 2.8

    # 导出内容不残留模板里的提示文字
    text = "\n".join(p.text for p in filled)
    assert "此处填写" not in text


def test_verify_blocks_uncited_and_unknown_refs(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    report_docs.set_outline(
        root,
        "RPT-0001",
        [{"id": "S1", "title": "一、基本情况", "level": 1}],
        approved=True,
    )
    report_docs.write_section(
        root,
        "RPT-0001",
        "S1",
        [
            {"type": "paragraph", "text": "无来源的一段话。"},
            {
                "type": "paragraph",
                "text": "引用了别处。",
                "sourceRefs": ["QZ-9999"],
            },
        ],
    )
    check = verify_report(root, "RPT-0001")
    codes = {error["code"] for error in check["errors"]}
    assert "E_UNCITED" in codes
    assert "E_UNKNOWN_SOURCE" in codes


def test_verify_requires_approved_outline_and_blocks(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    check = verify_report(root, "RPT-0001")
    codes = {error["code"] for error in check["errors"]}
    assert "E_OUTLINE_NOT_APPROVED" in codes
    assert "E_SECTION_EMPTY" in codes

    with pytest.raises(ApiError) as exc:
        report_docs.export_report(root, "RPT-0001")
    assert exc.value.code == "E_VERIFY_FAILED"


def test_verify_errors_on_untraceable_amount(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    report_docs.set_outline(
        root,
        "RPT-0001",
        [{"id": "S1", "title": "一、基本情况", "level": 1}],
        approved=True,
    )
    report_docs.write_section(
        root,
        "RPT-0001",
        "S1",
        [
            {
                "type": "paragraph",
                "text": "本次费用为8888元。",
                "sourceRefs": [form_id],
            }
        ],
    )
    check = verify_report(root, "RPT-0001")
    assert check["ok"] is False
    assert any(e["code"] == "E_AMOUNT_UNTRACED" for e in check["errors"])


def test_amount_tracing_handles_units_and_ignores_years(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root, amount="1.2")
    # 描述里写「1.2万元」，报告写「12000元」应视为同一笔；「2026元」不应被年份蒙混过关
    directory = evidence_forms.forms_dir(root) / form_id
    manifest = json.loads((directory / "manifest.json").read_text("utf-8"))
    manifest["findingsSnapshot"][0]["description"] = "采购费用1.2万元，发生于2026年。"
    (directory / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False), "utf-8"
    )
    created = report_docs.create_report(
        root,
        template_id=template_id,
        template_scope="project",
        form_ids=[form_id],
        header={"reportTitle": "金额测试"},
    )
    rid = created["manifest"]["id"]
    report_docs.set_outline(
        root,
        rid,
        [{"id": "S1", "title": "一、基本情况", "level": 1}],
        approved=True,
    )
    report_docs.write_section(
        root,
        rid,
        "S1",
        [
            {
                "type": "paragraph",
                "text": "采购费用12000元，发生于2026年。",
                "sourceRefs": [form_id],
            }
        ],
    )
    check = verify_report(root, rid)
    assert check["ok"] is True, check["errors"]

    report_docs.write_section(
        root,
        rid,
        "S1",
        [
            {
                "type": "paragraph",
                "text": "采购费用2026元。",
                "sourceRefs": [form_id],
            }
        ],
    )
    check = verify_report(root, rid)
    assert any(e["code"] == "E_AMOUNT_UNTRACED" for e in check["errors"])


def test_outline_ids_never_hang_on_collision(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    body = report_docs.set_outline(
        root,
        "RPT-0001",
        [
            {"title": "A", "id": "S3"},
            {"title": "B", "id": "S4"},
            {"title": "C"},
            {"title": "D", "id": "S3"},
        ],
        approved=True,
    )
    ids = [section["id"] for section in body["sections"]]
    assert len(ids) == len(set(ids)) == 4


def test_editing_after_export_marks_report_draft_again(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    _fill_all(root, form_id)
    report_docs.export_report(root, "RPT-0001")
    assert report_docs._load_manifest(root, "RPT-0001")["status"] == "saved"

    report_docs.write_section(
        root,
        "RPT-0001",
        "S1",
        [{"type": "paragraph", "text": "改一行。", "sourceRefs": [form_id]}],
    )
    assert report_docs._load_manifest(root, "RPT-0001")["status"] == "draft"


def test_export_preserves_leading_template_table(tmp_path):
    root, _template_id = _workspace(tmp_path)
    # 再注册一个带「红头表格」的模板
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(16)
    _east_asia(normal, "仿宋_GB2312")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "广西全灌高速公路有限公司文件"
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("审计报告")
    run.font.size = Pt(22)
    run.font.name = "小标宋"
    _east_asia(run, "小标宋")
    doc.add_paragraph("一、基本情况", style=doc.styles.add_style("标题 1", 1))
    doc.add_paragraph("占位。")

    item = config_store.create_item(
        "templates",
        {
            "scope": "project",
            "kind": "report",
            "name": "红头模板",
            "fileName": "leading.docx",
        },
        root,
    )
    target = templates_store.templates_root("project", root) / "report" / "leading.docx"
    doc.save(str(target))

    form_id = _form(root)
    created = report_docs.create_report(
        root,
        template_id=item["id"],
        template_scope="project",
        form_ids=[form_id],
        title="红头报告",
    )
    rid = created["manifest"]["id"]
    report_docs.set_outline(
        root, rid, [{"id": "S1", "title": "一、基本情况", "level": 1}], approved=True
    )
    report_docs.write_section(
        root,
        rid,
        "S1",
        [{"type": "paragraph", "text": "正文。", "sourceRefs": [form_id]}],
    )
    report_docs.export_report(root, rid)
    exported = Document(str(report_docs._dir(root, rid) / "document.docx"))
    assert exported.tables
    assert "广西全灌" in exported.tables[0].rows[0].cells[0].text


def test_export_applies_at_least_line_spacing(tmp_path):
    root, _template_id = _workspace(tmp_path)
    from docx.enum.text import WD_LINE_SPACING

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(16)
    _east_asia(normal, "仿宋_GB2312")
    normal.paragraph_format.line_spacing = Pt(28)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("审计报告")
    run.font.size = Pt(22)
    doc.add_paragraph("占位正文内容。")

    item = config_store.create_item(
        "templates",
        {
            "scope": "project",
            "kind": "report",
            "name": "最小行距模板",
            "fileName": "atleast.docx",
        },
        root,
    )
    target = templates_store.templates_root("project", root) / "report" / "atleast.docx"
    doc.save(str(target))

    form_id = _form(root)
    created = report_docs.create_report(
        root,
        template_id=item["id"],
        template_scope="project",
        form_ids=[form_id],
        title="行距报告",
    )
    rid = created["manifest"]["id"]
    report_docs.set_outline(
        root, rid, [{"id": "S1", "title": "一、基本情况", "level": 1}], approved=True
    )
    report_docs.write_section(
        root,
        rid,
        "S1",
        [{"type": "paragraph", "text": "正文。", "sourceRefs": [form_id]}],
    )
    report_docs.export_report(root, rid)
    exported = Document(str(report_docs._dir(root, rid) / "document.docx"))
    para = next(p for p in exported.paragraphs if p.text == "正文。")
    assert para.paragraph_format.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
    assert round(para.paragraph_format.line_spacing.pt) == 28


def test_user_origin_skips_citation_requirement(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    report_docs.set_outline(
        root,
        "RPT-0001",
        [{"id": "S1", "title": "一、基本情况", "level": 1}],
        approved=True,
    )
    report_docs.write_section(
        root,
        "RPT-0001",
        "S1",
        [
            {
                "type": "paragraph",
                "text": "本报告经部门负责人审阅。",
                "origin": "user",
            }
        ],
    )
    check = verify_report(root, "RPT-0001")
    assert check["ok"] is True


def test_chat_log_append_and_load(tmp_path):
    root, template_id = _workspace(tmp_path)
    form_id = _form(root)
    _create(root, template_id, [form_id])
    report_docs.append_chat(root, "RPT-0001", {"role": "user", "content": "写一版"})
    report_docs.append_chat(root, "RPT-0001", {"role": "assistant", "content": "好的"})
    entries = report_docs.load_chat(root, "RPT-0001")
    assert [entry["content"] for entry in entries] == ["写一版", "好的"]
